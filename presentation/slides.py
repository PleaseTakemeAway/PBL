"""Shared look of the talk decks: colors, fonts, shapes and the slide frames.

A deck script (presentation/build_<part>.py) builds its slides with these helpers, so all
parts look the same and a style change is made once. Run deck scripts from the repo root
with PYTHONPATH=. so `presentation.slides` imports.
"""

import datetime
import io
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

FIXED_TIME = datetime.datetime(2026, 1, 1)

FONT = "Malgun Gothic"  # 맑은 고딕: Korean font that ships with Office on Windows
MONO = "Consolas"

INK = "1F2A30"  # dark slate: title/closing background, body text
TEAL = "2E6F73"  # steel teal: structure, "done"
ORANGE = "E8743B"  # cutting-torch orange: model stages, highlights
GRAY = "5F6B73"
LIGHT = "EEF3F4"  # card tint
TEAL_TINT = "DCEBEC"
ORANGE_TINT = "FBE7DC"
WHITE = "FFFFFF"
CODE_BG = "263238"
CODE_FG = "E6EDF0"
PALE = "C9D3D6"  # light outline
DASH = 4  # MSO_LINE_DASH_STYLE.DASH
CHECK, CROSS, PART = "✓", "✕", "△"

W, H = 13.333, 7.5
X0, CW = 0.6, 12.133  # content left edge and width


# ---------------------------------------------------------------------- primitives
def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def style_run(run, size, color, bold=False, font=FONT, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):  # Korean glyphs use the East Asian font, not latin
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", font)


def set_bullet(paragraph, color):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", str(Inches(0.28)))
    ppr.set("indent", str(-Inches(0.22)))
    clr = ppr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": color})
    clr.append(srgb)
    ppr.append(clr)
    ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "•"}))


def write(
    tf,
    paras,
    *,
    size=16,
    color=INK,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    after=6,
    margin=0.0,
):
    """paras: list of paragraphs; a paragraph is a str, a run dict, or a list of runs.
    A dict paragraph may set bullet=True and its own size/color/bold/font/italic."""
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        runs = para if isinstance(para, list) else [para]
        bullet = False
        for r in runs:
            r = {"text": r} if isinstance(r, str) else r
            bullet = bullet or r.get("bullet", False)
            run = p.add_run()
            run.text = r["text"]
            style_run(
                run,
                r.get("size", size),
                r.get("color", color),
                r.get("bold", bold),
                r.get("font", font),
                r.get("italic", False),
            )
        if bullet:
            set_bullet(p, ORANGE)
    return tf


def text(slide, x, y, w, h, paras, **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    write(box.text_frame, paras, **kw)
    return box


def shape(slide, kind, x, y, w, h, fill, *, line=None, radius=0.06):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def card(slide, x, y, w, h, paras, *, fill=LIGHT, margin=0.18, **kw):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
    write(s.text_frame, paras, margin=margin, **kw)
    return s


def circle(slide, x, y, d, label, fill=ORANGE, color=WHITE, size=16):
    s = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)
    write(
        s.text_frame,
        [{"text": label, "bold": True}],
        size=size,
        color=color,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    return s


def arrow(slide, x, y, w=0.3, h=0.32, fill=ORANGE, kind=MSO_SHAPE.RIGHT_ARROW):
    return shape(slide, kind, x, y, w, h, fill)


def code(slide, x, y, w, h, lines, size=11):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, CODE_BG, radius=0.03)
    write(
        s.text_frame,
        [{"text": line} for line in lines],
        size=size,
        color=CODE_FG,
        font=MONO,
        after=2,
        margin=0.2,
    )
    return s


def chip(slide, x, y, w, h, label, *, fill=LIGHT, color=INK, size=13, bold=False):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=0.15)
    write(
        s.text_frame,
        [{"text": label, "bold": bold}],
        size=size,
        color=color,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        after=0,
        margin=0.05,
    )
    return s


# ---------------------------------------------------------------------- slide frames
class Deck:
    def __init__(self, out: Path, title: str, footer: str):
        self.out, self.title, self.footer = out, title, footer
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.count = 0

    def slide(self, notes, *, dark=False):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(INK if dark else WHITE)
        s.notes_slide.notes_text_frame.text = notes
        self.count += 1
        return s

    def content(self, title, message, notes):
        """A light slide with the title, a one-line message under it and the footer."""
        s = self.slide(notes)
        text(s, X0, 0.45, CW, 0.75, [title], size=34, bold=True, anchor=MSO_ANCHOR.TOP)
        text(s, X0, 1.15, CW, 0.45, [message], size=17, color=GRAY)
        text(s, X0, 6.95, 9.0, 0.3, [self.footer], size=10, color=GRAY)
        text(
            s,
            W - X0 - 1.0,
            6.95,
            1.0,
            0.3,
            [str(self.count)],
            size=10,
            color=GRAY,
            align=PP_ALIGN.RIGHT,
        )
        return s

    def save(self):
        props = self.prs.core_properties
        props.title = self.title
        props.author = "PBL cutplan_agent"
        props.created = props.modified = FIXED_TIME
        props.revision = 1
        buffer = io.BytesIO()
        self.prs.save(buffer)
        reproducible(buffer.getvalue(), self.out)
        return self.out


def reproducible(data: bytes, path: Path) -> None:
    """Rewrite the zip with fixed entry times, so the same slides give the same bytes and
    a rebuild without changes leaves git clean. A chart's embedded workbook carries the
    time it was written, so it is rewritten the same way."""
    with zipfile.ZipFile(io.BytesIO(data)) as src:
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                body = src.read(info.filename)
                if info.filename.endswith(".xlsx"):
                    body = fixed_workbook(body)
                entry = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
                entry.compress_type = zipfile.ZIP_DEFLATED
                dst.writestr(entry, body)


def fixed_workbook(data: bytes) -> bytes:
    """The workbook with its created and modified times set to FIXED_TIME."""
    stamp = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                body = src.read(info.filename)
                if info.filename == "docProps/core.xml":
                    body = re.sub(
                        rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*",
                        rb"\g<1>" + stamp,
                        body,
                    )
                entry = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
                entry.compress_type = zipfile.ZIP_DEFLATED
                dst.writestr(entry, body)
    return out.getvalue()


# ---------------------------------------------------------------------- charts and pictures
def plain_chart(chart):
    """No title and no frame around the chart area."""
    chart.has_title = False
    space = chart._chartSpace
    sp_pr = space.find(qn("c:spPr"))
    if sp_pr is None:
        sp_pr = space.makeelement(qn("c:spPr"), {})
        space.find(qn("c:chart")).addnext(sp_pr)
    sp_pr.append(sp_pr.makeelement(qn("a:noFill"), {}))
    line = sp_pr.makeelement(qn("a:ln"), {})
    line.append(line.makeelement(qn("a:noFill"), {}))
    sp_pr.append(line)


def bar_chart(
    slide,
    x,
    y,
    w,
    h,
    categories,
    values,
    colors,
    *,
    number_format="0",
    horizontal=True,
    label_size=14,
    gap=60,
):
    """One-series bar chart with a value label on each bar and no axis numbers or grid."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION

    data = CategoryChartData()
    data.categories = categories
    data.add_series("values", values)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = slide.shapes.add_chart(
        kind, Inches(x), Inches(y), Inches(w), Inches(h), data
    ).chart
    plain_chart(chart)
    chart.has_legend = False
    chart.font.size = Pt(label_size)
    chart.font.name = FONT
    chart.font.color.rgb = rgb(INK)
    plot = chart.plots[0]
    plot.gap_width = gap
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format = number_format
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(label_size)
    labels.font.bold = True
    for point, color in zip(plot.series[0].points, colors):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(color)
    value_axis = chart.value_axis
    value_axis.visible = False
    value_axis.has_major_gridlines = False
    value_axis.minimum_scale = 0
    category_axis = chart.category_axis
    category_axis.format.line.fill.background()
    category_axis.tick_labels.font.size = Pt(label_size)
    if horizontal:
        category_axis.reverse_order = True  # first category on top
    return chart


def doughnut_chart(slide, x, y, w, h, categories, values, colors, *, hole=58):
    """Doughnut with value labels and a legend on the right."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

    data = CategoryChartData()
    data.categories = categories
    data.add_series("values", values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), data
    ).chart
    plain_chart(chart)
    chart.font.size = Pt(14)
    chart.font.name = FONT
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(15)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(14)
    plot.data_labels.font.bold = True
    plot.data_labels.font.color.rgb = rgb(WHITE)
    plot.data_labels.number_format = "0"
    plot.data_labels.number_format_is_linked = False
    for point, color in zip(plot.series[0].points, colors):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = rgb(color)
    hole_size = plot._element.find(qn("c:holeSize"))
    if hole_size is None:
        hole_size = plot._element.makeelement(qn("c:holeSize"), {})
        plot._element.append(hole_size)
    hole_size.set("val", str(hole))
    return chart


def picture(slide, path, x, y, w=None, h=None):
    """A picture, or a gray box naming the missing file so the deck still builds."""
    if Path(path).is_file():
        return slide.shapes.add_picture(
            str(path),
            Inches(x),
            Inches(y),
            Inches(w) if w else None,
            Inches(h) if h else None,
        )
    return card(
        slide,
        x,
        y,
        w or 4,
        h or 3,
        [f"그림 없음: {path}"],
        size=12,
        color=GRAY,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def badge(slide, x, y, d, glyph, fill, color=WHITE, size=None):
    """A round marker with one symbol or number in it."""
    return circle(
        slide, x, y, d, glyph, fill=fill, color=color, size=size or int(d * 30)
    )


# ---------------------------------------------------------------------- composite shapes
def road(s, here: int, labels: list[str], w: float = 0.82):
    """The steps of a part in the top right corner, with the current step filled."""
    gap = 0.07
    x = W - X0 - (len(labels) * w + (len(labels) - 1) * gap)
    for i, label in enumerate(labels):
        chip(
            s,
            x + i * (w + gap),
            0.14,
            w,
            0.27,
            label,
            fill=ORANGE if i == here else LIGHT,
            color=WHITE if i == here else GRAY,
            size=10,
            bold=i == here,
        )


def outline(s, kind, x, y, w, h, color=GRAY, width=1.0):
    """A shape with a line and no fill."""
    sp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.background()
    sp.line.color.rgb = rgb(color)
    sp.line.width = Pt(width)
    sp.shadow.inherit = False
    return sp


def dashed(sp, color=ORANGE, width=1.5):
    sp.line.color.rgb = rgb(color)
    sp.line.width = Pt(width)
    sp.line.dash_style = DASH
    return sp


def connect(s, x1, y1, x2, y2, color=PALE, width=1.5):
    c = s.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    return c


def page(s, x, y, w, h, fill=WHITE, line_color=GRAY):
    """A sheet of paper (folded corner)."""
    return shape(s, MSO_SHAPE.FOLDED_CORNER, x, y, w, h, fill, line=line_color)


def real_sheet(s, x, y, w, h):
    """An abstract company drawing: frame, several parts, a parts table and a title block.
    Drawn with shapes so no customer drawing is shown."""
    shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, WHITE, line=GRAY)
    m = 0.05 * w
    outline(s, MSO_SHAPE.RECTANGLE, x + m, y + m, w - 2 * m, h - 2 * m, GRAY, 0.75)
    tw = 0.27 * w
    for r in range(6):  # parts table, top right
        outline(
            s,
            MSO_SHAPE.RECTANGLE,
            x + w - m - tw,
            y + m + r * 0.05 * h,
            tw,
            0.05 * h,
            PALE,
            0.5,
        )
    shape(  # title block, bottom right
        s,
        MSO_SHAPE.RECTANGLE,
        x + w - m - tw,
        y + h - m - 0.2 * h,
        tw,
        0.2 * h,
        LIGHT,
        line=PALE,
    )
    parts = [  # (kind, x, y, w, h) as fractions of the sheet
        (MSO_SHAPE.RECTANGLE, 0.09, 0.12, 0.28, 0.22),
        (MSO_SHAPE.OVAL, 0.10, 0.46, 0.13, 0.19),
        (MSO_SHAPE.OVAL, 0.29, 0.50, 0.09, 0.13),
        (MSO_SHAPE.RECTANGLE, 0.44, 0.10, 0.11, 0.30),
        (MSO_SHAPE.TRAPEZOID, 0.29, 0.72, 0.17, 0.12),
        (MSO_SHAPE.RECTANGLE, 0.50, 0.52, 0.13, 0.09),
        (MSO_SHAPE.OVAL, 0.55, 0.70, 0.07, 0.10),
        (MSO_SHAPE.RECTANGLE, 0.09, 0.76, 0.14, 0.06),
    ]
    for kind, px, py, pw, ph in parts:
        outline(s, kind, x + px * w, y + py * h, pw * w, ph * h, INK, 1.0)


def bubble(s, x, y, w, h, who, said, *, ai: bool):
    """One line of a conversation: who speaks and what."""
    fill, color = (INK, WHITE) if ai else (LIGHT, INK)
    said_run = {"text": said}
    if ai:
        said_run["font"] = MONO
    card(
        s,
        x,
        y,
        w,
        h,
        [[{"text": who + "  ", "bold": True, "size": 12, "color": ORANGE}, said_run]],
        fill=fill,
        color=color,
        size=15,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.15,
    )


def band(s, y, head, body, *, fill=ORANGE_TINT, color=ORANGE, h=0.85, size=17):
    """A full-width card with a bold lead word."""
    card(
        s,
        X0,
        y,
        CW,
        h,
        [[{"text": head + "  ", "bold": True, "color": color}, {"text": body}]],
        fill=fill,
        size=size,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )


def table(s, x, y, w, h, rows, widths, style):
    """A table; style(r, c) gives (fill, color, bold, size, align) for each cell."""
    t = s.shapes.add_table(
        len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)
    ).table
    for c, cw in enumerate(widths):
        t.columns[c].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            fill, color, bold, size, align = style(r, c)
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            write(
                cell.text_frame,
                [{"text": value}],
                size=size,
                bold=bold,
                color=color,
                align=align,
                after=0,
                margin=0.08,
            )
    return t


# ---------------------------------------------------------------------- slides
