"""Build presentation/roadmap.pptx: where the project stands, for listeners who are new to AI.

PYTHONPATH=. uv run --with python-pptx python presentation/build_roadmap.py

The story: v2 gave the whole job (drawings to purchase order) to Claude with 25 skills, and it
did not work well; v3 moved to local models on our own server and narrowed the job to reading
one drawing, tested on synthetic drawings (27/30); next comes harness engineering and tool
calling, so the local AI can use well-made tools on real drawings and, in the end, reach the
cutting plan and the purchase order. Few words on the slides, the narration is in the notes.

Where every number comes from: docs/presentation.md 5절. The cone picture is drawn at build
time from data_v3/synthetic (presentation/build_dwg_dxf.py); real drawings appear only as
counts and as an abstract sketch.
"""

from pathlib import Path

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from presentation.build_dwg_dxf import render_cone
from presentation.slides import (
    CW,
    GRAY,
    INK,
    LIGHT,
    MONO,
    ORANGE,
    ORANGE_TINT,
    TEAL,
    TEAL_TINT,
    WHITE,
    X0,
    Deck,
    W,
    arrow,
    badge,
    bar_chart,
    card,
    chip,
    circle,
    code,
    picture,
    rgb,
    shape,
    text,
    write,
)

OUT = Path(__file__).with_name("roadmap.pptx")
TITLE = "도면에서 커팅플랜까지: 우리는 지금 어디쯤 왔나"
FOOTER = "PBL cutplan_agent · 로드맵 · 2026-10-24"
CHECK, CROSS, PART = "✓", "✕", "△"
PALE = "C9D3D6"
DASH = 4  # MSO_LINE_DASH_STYLE.DASH
ROAD = ["v2", "v3", "지금", "다음"]

# runs/reader/deepseek_ocr_qwen_reasoning_skills/accuracy.md (colleague, 2026-10-18)
SHAPES = [("원뿔", 6), ("프레스 성형", 6), ("절곡판", 6), ("평판", 5), ("원통", 4)]

# The v3 stages, as in docs/reader/quickstart.md; True = a model answers there
STAGES = [
    ("01", "도면 읽기", False),
    ("02", "표제란 찾기", False),
    ("03", "표제란 받아쓰기", True),
    ("04", "표제란 정리", True),
    ("05", "뷰 나누기", False),
    ("06", "뷰 · 모양 판단", True),
    ("07", "합치기", False),
    ("08", "3D 용 정리", False),
    ("09", "전개 치수", True),
    ("10", "전개 계산", False),
    ("11", "3D 후보", False),
    ("12", "3D 후보 고르기", True),
    ("13", "저장", False),
]


# ---------------------------------------------------------------------- local shapes
def road(s, here: int):
    """The four-step road in the top right corner, with the current step filled."""
    w, gap = 0.82, 0.07
    x = W - X0 - (4 * w + 3 * gap)
    for i, label in enumerate(ROAD):
        chip(
            s,
            x + i * (w + gap),
            0.14,
            w,
            0.27,
            label,
            fill=ORANGE if i == here else LIGHT,
            color=WHITE if i == here else GRAY,
            size=10,
            bold=i == here,
        )


def outline(s, kind, x, y, w, h, color=GRAY, width=1.0):
    """A shape with a line and no fill."""
    sp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.background()
    sp.line.color.rgb = rgb(color)
    sp.line.width = Pt(width)
    sp.shadow.inherit = False
    return sp


def dashed(sp, color=ORANGE, width=1.5):
    sp.line.color.rgb = rgb(color)
    sp.line.width = Pt(width)
    sp.line.dash_style = DASH
    return sp


def connect(s, x1, y1, x2, y2, color=PALE, width=1.5):
    c = s.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    return c


def page(s, x, y, w, h, fill=WHITE, line_color=GRAY):
    """A sheet of paper (folded corner)."""
    return shape(s, MSO_SHAPE.FOLDED_CORNER, x, y, w, h, fill, line=line_color)


def real_sheet(s, x, y, w, h):
    """An abstract company drawing: frame, several parts, a parts table and a title block.
    Drawn with shapes so no customer drawing is shown."""
    shape(s, MSO_SHAPE.RECTANGLE, x, y, w, h, WHITE, line=GRAY)
    m = 0.05 * w
    outline(s, MSO_SHAPE.RECTANGLE, x + m, y + m, w - 2 * m, h - 2 * m, GRAY, 0.75)
    tw = 0.27 * w
    for r in range(6):  # parts table, top right
        outline(
            s,
            MSO_SHAPE.RECTANGLE,
            x + w - m - tw,
            y + m + r * 0.05 * h,
            tw,
            0.05 * h,
            PALE,
            0.5,
        )
    shape(  # title block, bottom right
        s,
        MSO_SHAPE.RECTANGLE,
        x + w - m - tw,
        y + h - m - 0.2 * h,
        tw,
        0.2 * h,
        LIGHT,
        line=PALE,
    )
    parts = [  # (kind, x, y, w, h) as fractions of the sheet
        (MSO_SHAPE.RECTANGLE, 0.09, 0.12, 0.28, 0.22),
        (MSO_SHAPE.OVAL, 0.10, 0.46, 0.13, 0.19),
        (MSO_SHAPE.OVAL, 0.29, 0.50, 0.09, 0.13),
        (MSO_SHAPE.RECTANGLE, 0.44, 0.10, 0.11, 0.30),
        (MSO_SHAPE.TRAPEZOID, 0.29, 0.72, 0.17, 0.12),
        (MSO_SHAPE.RECTANGLE, 0.50, 0.52, 0.13, 0.09),
        (MSO_SHAPE.OVAL, 0.55, 0.70, 0.07, 0.10),
        (MSO_SHAPE.RECTANGLE, 0.09, 0.76, 0.14, 0.06),
    ]
    for kind, px, py, pw, ph in parts:
        outline(s, kind, x + px * w, y + py * h, pw * w, ph * h, INK, 1.0)


def bubble(s, x, y, w, h, who, said, *, ai: bool):
    """One line of a conversation: who speaks and what."""
    fill, color = (INK, WHITE) if ai else (LIGHT, INK)
    said_run = {"text": said}
    if ai:
        said_run["font"] = MONO
    card(
        s,
        x,
        y,
        w,
        h,
        [[{"text": who + "  ", "bold": True, "size": 12, "color": ORANGE}, said_run]],
        fill=fill,
        color=color,
        size=15,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.15,
    )


def band(s, y, head, body, *, fill=ORANGE_TINT, color=ORANGE, h=0.85, size=17):
    """A full-width card with a bold lead word."""
    card(
        s,
        X0,
        y,
        CW,
        h,
        [[{"text": head + "  ", "bold": True, "color": color}, {"text": body}]],
        fill=fill,
        size=size,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )


def table(s, x, y, w, h, rows, widths, style):
    """A table; style(r, c) gives (fill, color, bold, size, align) for each cell."""
    t = s.shapes.add_table(
        len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)
    ).table
    for c, cw in enumerate(widths):
        t.columns[c].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            fill, color, bold, size, align = style(r, c)
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            write(
                cell.text_frame,
                [{"text": value}],
                size=size,
                bold=bold,
                color=color,
                align=align,
                after=0,
                margin=0.08,
            )
    return t


# ---------------------------------------------------------------------- slides
def title_slide(d: Deck):
    s = d.slide(
        "이번 파트에서는 이 프로젝트가 지금 어디쯤 와 있는지를 말씀드리겠습니다. 처음에는 Claude 라는 "
        "AI 에게 일을 통째로 맡겨 봤고(v2), 그다음에는 우리 서버의 AI 로 도면 읽기에만 집중했습니다(v3). "
        "그 결과와, 앞으로 무엇을 할지까지 순서대로 가겠습니다.",
        dark=True,
    )
    text(
        s,
        X0,
        1.9,
        CW,
        1.6,
        ["도면에서 커팅플랜까지,", "우리는 지금 어디쯤 왔나"],
        size=40,
        bold=True,
        color=WHITE,
        after=4,
    )
    text(
        s,
        X0,
        3.75,
        CW,
        0.5,
        ["v2 → v3 → 다음: 로컬 AI 와 tool calling"],
        size=20,
        color=ORANGE,
    )
    w, gap = 2.4, 0.6
    for i, label in enumerate(ROAD):
        x = X0 + i * (w + gap)
        chip(
            s,
            x,
            5.0,
            w,
            0.55,
            label,
            fill=ORANGE if label == "지금" else "34444C",
            color=WHITE,
            size=16,
            bold=True,
        )
        if i < 3:
            arrow(s, x + w + 0.15, 5.12, 0.3, 0.3, fill="5F6B73")
    text(
        s,
        X0,
        6.4,
        CW,
        0.4,
        ["PBL cutplan_agent · 2026-10-24"],
        size=14,
        color="B8C4C9",
    )


def map_slide(d: Deck):
    s = d.content(
        "전체 지도: 도면에서 구매의뢰까지",
        "회사 일 네 단계 중 AI 가 어디까지 맡았나",
        "먼저 전체 지도입니다. 회사에서 하는 일은 네 단계입니다. 도면을 읽어 부품의 모양과 치수, 재질을 "
        "알아내고, 부품을 평평하게 펼친 전개도를 그리고, 전개도 조각들을 철판 위에 배치하는 커팅플랜을 "
        "짜고, 필요한 철판을 주문하는 구매의뢰서를 씁니다. v2 에서는 이 네 단계를 전부 AI 에게 맡겼는데 "
        "잘 안 됐습니다. v3 에서는 앞의 두 단계, 그중에서도 도면 읽기에 집중했고, 연습용 도면에서 꽤 "
        "좋은 결과를 얻었습니다. 지금 우리는 여기에 있습니다. 다음 목표는 실제 회사 도면으로, 그리고 "
        "끝까지 가는 것입니다.",
    )
    road(s, 2)
    lw, gap = 1.55, 0.3
    cx = X0 + lw + 0.15
    cw = (CW - lw - 0.15 - 3 * gap) / 4
    steps = [
        ("① 도면 읽기", "모양 · 치수 · 재질"),
        ("② 전개도", "펼친 모양 계산"),
        ("③ 커팅플랜", "철판 위에 배치"),
        ("④ 구매의뢰", "철판 주문서"),
    ]
    for i, (head, sub) in enumerate(steps):
        x = cx + i * (cw + gap)
        box = card(
            s,
            x,
            1.85,
            cw,
            1.0,
            [
                {"text": head, "bold": True, "size": 19},
                {"text": sub, "size": 14, "color": GRAY},
            ],
            fill=WHITE,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            after=2,
        )
        box.line.color.rgb = rgb(PALE)
        if i < 3:
            arrow(s, x + cw + 0.04, 2.2, 0.22, 0.3)
    rows = [
        ("v2", "Claude", 4, LIGHT, GRAY, "AI 에게 전부 맡김 → 잘 안 됨", CROSS),
        ("v3", "로컬 AI", 2, TEAL, WHITE, "도면 읽기 → 연습 도면 27/30", CHECK),
        (
            "다음",
            "로컬 AI + 도구",
            4,
            ORANGE_TINT,
            ORANGE,
            "실제 도면으로, 끝까지",
            None,
        ),
    ]
    for r, (name, who, span, fill, color, label, mark) in enumerate(rows):
        y = 3.15 + r * 1.12
        text(
            s,
            X0,
            y,
            lw,
            0.85,
            [
                {"text": name, "bold": True, "size": 20},
                {"text": who, "size": 13, "color": GRAY},
            ],
            anchor=MSO_ANCHOR.MIDDLE,
            after=0,
        )
        bw = span * cw + (span - 1) * gap
        bar = card(
            s,
            cx,
            y + 0.08,
            bw,
            0.7,
            [{"text": label, "bold": True}],
            fill=fill,
            color=color,
            size=17,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.25,
        )
        if name == "다음":
            dashed(bar)
        if mark:
            badge(
                s,
                cx + bw - 0.62,
                y + 0.16,
                0.54,
                mark,
                "AAB6BB" if mark == CROSS else WHITE,
                color=WHITE if mark == CROSS else TEAL,
                size=18,
            )
        if name == "v3":
            chip(
                s,
                cx + bw + 0.25,
                y + 0.16,
                1.75,
                0.54,
                "◀ 지금 여기",
                fill=ORANGE,
                color=WHITE,
                size=16,
                bold=True,
            )


def v2_slide(d: Deck):
    s = d.content(
        "v2: AI 에게 일 전체를 맡겼다",
        "가장 강한 Claude 가 도구 25개를 골라 쓰며 도면부터 구매의뢰까지",
        "v2 는 이렇게 생겼습니다. 왼쪽의 도면 세트를 Claude 에게 줍니다. Claude 는 Anthropic 이라는 회사의 "
        "AI 이고, 그중에서도 가장 강한 모델(Opus)을 썼습니다. Claude 옆에는 도구 25개가 있습니다. 도면을 "
        "열고 글자를 읽는 도구, 도면 종류나 부품 모양을 판단하는 도구, 전개와 네스팅과 구매의뢰를 "
        "계산하는 도구, 무게나 겹침을 검사하고 사람에게 묻는 도구입니다. 이 도구를 skill 이라고 "
        "불렀습니다. Claude 는 이 도구를 스스로 골라 쓰면서 도면에서 구매의뢰서까지 가야 했습니다. "
        "도면을 보는 것도 Claude, 순서를 정하는 것도 Claude 였습니다. 대상은 연습용 도면이 아니라 실제 "
        "회사 도면(압력용기, 판금 부품)이었습니다.",
    )
    road(s, 0)
    for k in range(3):
        page(s, X0 + 0.18 * k, 2.3 + 0.18 * k, 1.3, 1.75)
    text(
        s,
        X0 - 0.1,
        4.45,
        1.9,
        0.7,
        [
            {"text": "도면 세트", "bold": True, "size": 16},
            {"text": "실제 회사 도면", "size": 13, "color": GRAY},
        ],
        align=PP_ALIGN.CENTER,
        after=0,
    )
    arrow(s, 2.5, 3.35, 0.35, 0.36)
    circle(s, 3.0, 2.45, 2.0, "Claude", fill=INK, size=24)
    text(
        s,
        2.6,
        4.6,
        2.8,
        0.7,
        ["보는 것도, 순서를 정하는 것도"],
        size=13,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )
    arrow(s, 5.15, 3.3, 0.5, 0.42, kind=MSO_SHAPE.LEFT_RIGHT_ARROW)
    groups = [
        ("4", "보기 · 읽기", "도면 열기, 글자 읽기"),
        ("7", "판단", "도면 종류, 부품 모양"),
        ("6", "계산", "전개, 네스팅, 구매의뢰"),
        ("6", "검사", "무게 · 겹침, 사람 확인"),
    ]
    gx, gw, gh = 5.8, 2.45, 1.3
    for i, (n, head, sub) in enumerate(groups):
        r, c = divmod(i, 2)
        card(
            s,
            gx + c * (gw + 0.15),
            2.05 + r * (gh + 0.15),
            gw,
            gh,
            [
                [
                    {"text": n, "bold": True, "size": 30, "color": ORANGE},
                    {"text": "  " + head, "bold": True, "size": 18},
                ],
                {"text": sub, "size": 13, "color": GRAY},
            ],
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.2,
            after=2,
        )
    text(
        s,
        gx,
        4.85,
        2 * gw + 0.15,
        0.4,
        ["+ 상태 확인 2개 = 도구(skill) 25개"],
        size=14,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )
    arrow(s, 10.95, 3.35, 0.35, 0.36)
    page(s, 11.4, 2.4, 1.33, 1.75, fill=ORANGE_TINT, line_color=ORANGE)
    text(
        s,
        11.1,
        4.25,
        1.9,
        0.7,
        [
            {"text": "커팅플랜", "bold": True, "size": 15},
            {"text": "구매의뢰서", "bold": True, "size": 15},
        ],
        align=PP_ALIGN.CENTER,
        after=0,
    )
    band(
        s,
        5.55,
        "v2",
        "Claude Opus (클라우드 AI) 하나가 ① 도면 읽기부터 ④ 구매의뢰까지 전부",
        fill=LIGHT,
        color=INK,
    )


def skill_slide(d: Deck):
    s = d.content(
        "skill 은 무엇인가",
        "프로그램이 할 줄 아는 일 하나 + AI 가 읽는 이름과 설명",
        "skill 이 무엇인지 식당에 비유해 보겠습니다. AI 는 손님이고, 프로그램은 주방입니다. 손님은 "
        "메뉴판을 봅니다. 메뉴판에는 요리 이름과 설명, 주문할 때 적을 것이 있습니다. skill 도 똑같이 "
        "이름, 설명, 입력 양식이 있습니다. AI 가 'unfold_part 를 부품 3번으로'라고 주문하면, 실제 계산은 "
        "주방, 즉 우리가 미리 짜 둔 프로그램이 합니다. AI 가 직접 원주율을 곱하는 게 아닙니다. 그리고 "
        "skill 에는 참고서도 있습니다. 전개 규칙, 도면 읽는 법 같은 도메인 지식 문서 7권인데, AI 가 "
        "필요할 때 펼쳐 봅니다.",
    )
    road(s, 0)
    card(
        s,
        X0,
        1.85,
        5.0,
        3.65,
        [{"text": "메뉴판 — AI 가 보는 것", "bold": True, "size": 18, "color": TEAL}],
        fill=TEAL_TINT,
        margin=0.25,
    )
    menu = [
        ("unfold_part", "부품을 펼친 크기를 계산한다", "부품 번호"),
        ("nest_all", "조각을 철판 위에 배치한다", "자르는 방식"),
        ("ask_human", "사람에게 확인을 받는다", "질문"),
    ]
    for i, (name, desc, form) in enumerate(menu):
        card(
            s,
            X0 + 0.2,
            2.62 + i * 0.93,
            4.6,
            0.8,
            [
                [
                    {"text": name, "bold": True, "font": MONO, "size": 15},
                    {"text": "   입력: " + form, "size": 12, "color": GRAY},
                ],
                {"text": desc, "size": 14},
            ],
            fill=WHITE,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.15,
            after=1,
        )
    circle(s, 6.05, 2.2, 1.1, "AI", fill=INK, size=22)
    card(
        s,
        5.75,
        3.5,
        1.7,
        0.95,
        [
            {"text": "주문", "bold": True, "size": 13, "color": ORANGE},
            {"text": "unfold_part(3)", "size": 12, "font": MONO, "bold": True},
        ],
        fill=ORANGE_TINT,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.05,
        after=0,
    )
    arrow(s, 7.5, 3.78, 0.32, 0.36)
    kx = 7.95
    kw = CW - (kx - X0)
    text(
        s,
        kx,
        1.85,
        kw,
        0.5,
        [{"text": "주방 — 프로그램이 하는 일", "bold": True}],
        size=18,
        color=ORANGE,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    code(
        s,
        kx,
        2.45,
        kw,
        3.05,
        [
            "def unfold_part(part_index):",
            '    """부품을 펼친 크기를 계산한다"""',
            "    p = 부품[part_index]",
            "    if p.모양 == '원통':",
            "        길이 = π × (지름 + 두께)",
            "    ...",
            "    return {'전개 길이': 길이, ...}",
        ],
        size=14,
    )
    band(
        s,
        5.75,
        "skill 문서 = 참고서 7권",
        "전개 규칙, 도면 읽는 법, 네스팅 규칙 … — AI 가 필요할 때 펼쳐 본다",
        fill=LIGHT,
        color=TEAL,
    )


def claude_loop_slide(d: Deck):
    s = d.content(
        "Claude 는 skill 을 이렇게 썼다",
        "Claude 는 주문만 하고 일은 프로그램이 한다. 이것을 수십 번 되풀이",
        "Claude 가 skill 을 쓰는 모습을 대화로 옮겨 봤습니다. 프로그램이 먼저 도면 목록과 메뉴판 25개를 "
        "줍니다. Claude 는 '도면 1번을 열어 줘'라고 주문하고, 프로그램이 실행해서 결과를 돌려줍니다. "
        "Claude 가 그 결과를 보고 다음 주문을 합니다. 도면 종류를 판단하고, 자재표를 읽고, 전개하고, "
        "배치하고, 구매의뢰서를 쓸 때까지 이걸 수십 번 반복합니다. 오른쪽이 그 반복을 그림으로 그린 "
        "것입니다. 이 방식에는 이름이 있는데, 뒤에서 다시 나옵니다.",
    )
    road(s, 0)
    lines = [
        ("프로그램", "도면 12장과 메뉴판(skill 25개)을 줄게", False),
        ("Claude", "ingest_drawing(도면 1)", True),
        ("프로그램", "준비 끝: 그림 3쪽", False),
        ("Claude", "classify_drawing(도면 1)", True),
        ("프로그램", "본체 상세도, 자재표 있음", False),
        ("Claude", "… nest_all() → build_por()", True),
        ("프로그램", "커팅플랜 · 구매의뢰서 완성", False),
    ]
    bw = 6.3
    for i, (who, said, ai) in enumerate(lines):
        bubble(s, X0 + (1.2 if ai else 0), 1.85 + i * 0.68, bw, 0.56, who, said, ai=ai)
    lx = X0 + 7.95
    lw = CW - 7.95
    loop = [
        ("1", "메뉴판을 보고 고른다", "Claude"),
        ("2", "주문한다 — 이름 + 입력값", "Claude"),
        ("3", "실행한다", "프로그램"),
        ("4", "결과를 돌려준다", "프로그램"),
    ]
    for i, (num, head, who) in enumerate(loop):
        y = 1.95 + i * 1.0
        circle(s, lx, y, 0.62, num, fill=INK if who == "Claude" else TEAL, size=16)
        text(
            s,
            lx + 0.8,
            y - 0.05,
            lw - 0.8,
            0.72,
            [
                {"text": head, "bold": True, "size": 16},
                {"text": who, "size": 13, "color": GRAY},
            ],
            anchor=MSO_ANCHOR.MIDDLE,
            after=0,
        )
    arrow(s, lx + 0.16, 5.95, 0.3, 0.36, kind=MSO_SHAPE.UP_ARROW)
    text(
        s,
        lx + 0.6,
        5.9,
        lw - 0.6,
        0.5,
        ["다시 1 로 — 끝날 때까지"],
        size=15,
        color=ORANGE,
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def v2_result_slide(d: Deck):
    s = d.content(
        "v2 결과: 잘 안 됐다",
        "일이 너무 컸다. 첫 단추(도면 읽기)가 어긋나면 뒤가 전부 틀린다",
        "결과는 잘 안 됐습니다. 이유는 일이 너무 컸다는 데 있습니다. 네 단계가 줄줄이 이어져 있어서, 첫 "
        "단계인 도면 읽기에서 숫자 하나가 틀리면 그 뒤의 전개도, 커팅플랜, 구매의뢰가 전부 틀립니다. 예를 "
        "들어 원통의 지름을 10 mm 잘못 읽으면 펼친 길이는 31 mm 가 달라지고, v2 의 AI 지침서에 적어 둔 "
        "현장 기준으로 그렇게 자른 철판은 못 씁니다. 그리고 단계마다 맞았는지 재지 않으면 어디서 "
        "틀렸는지 찾기 어렵습니다. 그래서 배운 것은, "
        "첫 단추인 도면 읽기부터 제대로 해야 한다는 것이었습니다.",
    )
    road(s, 0)
    steps = ["① 도면 읽기", "② 전개도", "③ 커팅플랜", "④ 구매의뢰"]
    w, gap = 2.6, 0.58
    for i, name in enumerate(steps):
        x = X0 + i * (w + gap)
        first = i == 0
        card(
            s,
            x,
            1.95,
            w,
            1.15,
            [{"text": name, "bold": True, "size": 20}],
            fill=ORANGE_TINT if first else LIGHT,
            color=INK if first else GRAY,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        badge(
            s, x + w - 0.45, 1.75, 0.55, CROSS, ORANGE if first else "AAB6BB", size=18
        )
        if i < 3:
            arrow(s, x + w + 0.13, 2.36, 0.32, 0.34, fill=ORANGE if first else PALE)
    text(
        s,
        X0,
        3.2,
        w,
        0.4,
        ["숫자 하나가 어긋나면"],
        size=14,
        color=ORANGE,
        align=PP_ALIGN.CENTER,
    )
    text(
        s,
        X0 + w + gap,
        3.2,
        3 * w + 2 * gap,
        0.4,
        ["뒤는 전부 틀린다"],
        size=14,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )
    chain = [
        ("10 mm", "예: 지름을 잘못 읽으면"),
        ("31 mm", "펼친 길이가 달라지고"),
        ("폐기", "그렇게 자른 철판은"),
    ]
    cw2 = 3.5
    gap2 = (CW - 3 * cw2) / 2
    for i, (big, small) in enumerate(chain):
        x = X0 + i * (cw2 + gap2)
        box = card(
            s,
            x,
            3.8,
            cw2,
            1.35,
            [
                {"text": small, "size": 15, "color": GRAY},
                {"text": big, "bold": True, "size": 36, "color": ORANGE},
            ],
            fill=WHITE,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            after=0,
        )
        box.line.color.rgb = rgb(PALE)
        if i < 2:
            arrow(s, x + cw2 + gap2 / 2 - 0.18, 4.3, 0.36, 0.36)
    band(
        s,
        5.65,
        "배운 것",
        "단계마다 맞았는지 재지 않으면 어디서 틀렸는지 찾기 어렵다 → 첫 단추부터 제대로",
    )
    text(
        s,
        X0,
        5.22,
        CW,
        0.3,
        [
            "예시의 기준: v2 의 AI 지침서에 적힌 현장 기준 (지름 10 mm 오차 = 둘레 31 mm, 그 판은 폐기)"
        ],
        size=11,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )


def why_move_slide(d: Deck):
    s = d.content(
        "왜 옮겼나",
        "수업의 목표, v2 의 교훈, 비용과 보안",
        "v3 로 넘어간 이유는 세 가지입니다. 첫째, 이 수업의 목표입니다. 비싼 AI 하나가 아니라 값싼 로컬 "
        "AI 로도, 주변 장치인 하네스를 잘 만들면 쓸 만한 결과를 낼 수 있다는 걸 보이는 것이 목표입니다. "
        "Claude 는 처음에 무엇이 필요한지 찾는 시행착오용이었습니다. 둘째, v2 의 교훈입니다. 일이 너무 "
        "컸습니다. 더 작은 로컬 AI 에게는 더더욱 그렇습니다. 그래서 첫 단추인 도면 읽기부터, 좁은 질문으로 "
        "시작했습니다. 셋째, 비용과 보안입니다. 클라우드 AI 는 부를 때마다 돈이 들고, 회사 도면이 바깥 "
        "서버로 나갑니다. 로컬 AI 는 우리 서버 안에서 돕니다.",
    )
    road(s, 1)
    reasons = [
        (
            "1",
            "수업의 목표",
            "값싼 로컬 AI + 하네스로도 된다는 것을 보인다",
            "Claude 는 무엇이 필요한지 찾는 시행착오용",
        ),
        (
            "2",
            "v2 의 교훈",
            "일이 너무 컸다. 작은 AI 에게는 더더욱",
            "→ 첫 단추(도면 읽기)부터, 좁은 질문으로",
        ),
        (
            "3",
            "비용과 보안",
            "클라우드 AI 는 부를 때마다 비용, 도면이 바깥 서버로",
            "→ 로컬 AI 는 우리 서버 안에서",
        ),
    ]
    w, gap = 3.85, 0.29
    for i, (num, head, body, sub) in enumerate(reasons):
        x = X0 + i * (w + gap)
        card(
            s,
            x,
            1.95,
            w,
            3.5,
            [
                {"text": head, "bold": True, "size": 24, "color": TEAL},
                {"text": body, "bold": True, "size": 17},
                {"text": sub, "size": 15, "color": GRAY},
            ],
            margin=0.3,
            after=10,
        ).text_frame.margin_top = Inches(1.3)
        circle(s, x + 0.3, 2.25, 0.8, num, size=24)
    band(s, 5.7, "그래서 v3", "누가 → 우리 서버의 로컬 AI,  무엇을 → 도면 읽기 하나")


def v3_shift_slide(d: Deck):
    s = d.content(
        "v3: 무게중심을 옮겼다",
        "누가 → 우리 서버의 로컬 AI, 무엇을 → 도면 읽기 하나",
        "그래서 무엇을 옮겼는지 정리하면 이렇습니다. 첫째, 누가 하느냐입니다. Claude 는 회사 밖 클라우드에 "
        "있어서 부를 때마다 돈이 들고 도면이 밖으로 나갑니다. v3 는 우리 서버에 직접 띄운 로컬 AI 를 "
        "씁니다. 둘째, 무엇을 하느냐입니다. 네 단계 전부가 아니라 첫 단추인 도면 읽기 하나에 "
        "집중했습니다. 전개도는 읽은 치수로 프로그램이 계산합니다. 셋째, 순서는 AI 가 정하지 않고 코드에 "
        "고정했습니다. AI 는 정해진 다섯 곳에서 질문에 답만 합니다.",
    )
    road(s, 1)
    rows = [
        (
            "누가",
            ("Claude — 회사 밖 클라우드", "부를 때마다 비용 · 도면이 밖으로"),
            ("로컬 AI — 우리 서버", "추가 비용 없음 · 도면이 안 나간다"),
        ),
        (
            "무엇을",
            ("① → ④ 전부", "도면부터 구매의뢰서까지"),
            ("① 도면 읽기에 집중", "② 전개도는 읽은 치수로 프로그램이 계산"),
        ),
        (
            "어떻게",
            ("AI 가 순서까지 정한다", "도구 25개 중 골라 쓰기"),
            ("순서는 코드가 정한다", "AI 는 정해진 다섯 곳에서 답만"),
        ),
    ]
    lw, bw, aw, agap = 1.5, 4.4, 0.55, 0.2
    bx = X0 + lw + 0.2
    ax = bx + bw + agap
    afx = ax + aw + agap
    afw = CW - (afx - X0)
    for i, (label, before, after) in enumerate(rows):
        y = 1.95 + i * 1.5
        chip(s, X0, y + 0.3, lw, 0.65, label, fill=INK, color=WHITE, size=18, bold=True)
        for x, w, (head, sub), fill, color in (
            (bx, bw, before, LIGHT, GRAY),
            (afx, afw, after, TEAL_TINT, TEAL),
        ):
            card(
                s,
                x,
                y,
                w,
                1.25,
                [
                    {"text": head, "bold": True, "size": 19, "color": color},
                    {"text": sub, "size": 14, "color": GRAY},
                ],
                fill=fill,
                anchor=MSO_ANCHOR.MIDDLE,
                margin=0.25,
                after=3,
            )
        arrow(s, ax, y + 0.43, aw, 0.4)
    text(s, bx, 6.45, bw, 0.35, ["v2"], size=13, color=GRAY, bold=True)
    text(s, afx, 6.45, afw, 0.35, ["v3"], size=13, color=TEAL, bold=True)


def server_slide(d: Deck):
    s = d.content(
        "로컬 AI 는 이렇게 띄웠다",
        "GPU 4장짜리 서버 한 대에 AI 두 개: 글자 받아쓰기 + 그림 보고 판단",
        "로컬 AI 를 어떻게 띄웠는지 보겠습니다. AI 모델은 숫자가 가득 든 큰 파일이라고 생각하시면 됩니다. "
        "이 파일을 GPU 라는 계산 장치의 메모리에 올려야 돌아갑니다. 서버에는 메모리 24 GB 짜리 GPU 가 "
        "4장 있습니다. 첫 번째 장에는 글자 받아쓰기 전문인 OCR 모델을 올렸습니다. 6 GB 로 작습니다. "
        "나머지 세 장에는 그림과 글자를 같이 보고 판단하는 추론 모델 Qwen 을 올렸습니다. 52 GB 라서 한 "
        "장에 안 들어가, 세 장에 나눠 올렸습니다. 클수록 똑똑하지만 그만큼 GPU 가 더 필요합니다. 참고로 "
        "동료는 90 GB 급 GPU 두 장으로 같은 실험을 했습니다.",
    )
    road(s, 1)
    sw = 7.25
    card(
        s,
        X0,
        1.85,
        sw,
        3.55,
        [{"text": "작업 서버 1대", "bold": True, "size": 17}],
        fill=LIGHT,
        margin=0.25,
    )
    gw, gg = 1.58, 0.13
    gx0 = X0 + 0.3
    for i in range(4):
        card(
            s,
            gx0 + i * (gw + gg),
            2.5,
            gw,
            1.55,
            [
                {"text": f"GPU {i}", "bold": True, "size": 17},
                {"text": "24 GB", "size": 14},
            ],
            fill=TEAL if i == 0 else ORANGE,
            color=WHITE,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            after=2,
        )
    shape(s, MSO_SHAPE.RECTANGLE, gx0, 4.18, gw, 0.07, TEAL)
    shape(s, MSO_SHAPE.RECTANGLE, gx0 + gw + gg, 4.18, 3 * gw + 2 * gg, 0.07, ORANGE)
    text(
        s,
        gx0,
        4.3,
        gw,
        0.9,
        [{"text": "OCR 모델", "bold": True}, {"text": "6 GB", "color": GRAY}],
        size=15,
        align=PP_ALIGN.CENTER,
        after=0,
    )
    text(
        s,
        gx0 + gw + gg,
        4.3,
        3 * gw + 2 * gg,
        0.9,
        [
            {"text": "추론 모델 — 3장에 나눠 올림", "bold": True},
            {"text": "52 GB (한 장 24 GB 에 안 들어간다)", "color": GRAY},
        ],
        size=15,
        align=PP_ALIGN.CENTER,
        after=0,
    )
    x = X0
    for spec, w in zip(
        ["GPU  RTX A5000 24 GB × 4", "CPU  Xeon 48코어", "메모리  128 GB"],
        [3.1, 2.0, 1.85],
    ):
        chip(
            s, x, 5.6, w, 0.5, spec, fill=WHITE, size=13, bold=True
        ).line.color.rgb = rgb(PALE)
        x += w + 0.15
    rx = X0 + sw + 0.3
    rw = CW - sw - 0.3
    roles = [
        (
            "글자 받아쓰기",
            "OCR 모델 · DeepSeek-OCR-2",
            "표제란 그림 → 글자",
            TEAL,
            TEAL_TINT,
        ),
        (
            "그림 보고 판단",
            "추론 모델 (VLM) · Qwen3.6-27B",
            "그림 + 글자 → 정해진 양식의 답",
            ORANGE,
            ORANGE_TINT,
        ),
    ]
    for i, (head, model, job, color, fill) in enumerate(roles):
        card(
            s,
            rx,
            1.85 + i * 1.9,
            rw,
            1.7,
            [
                {"text": head, "bold": True, "size": 21, "color": color},
                {"text": model, "size": 14, "bold": True},
                {"text": job, "size": 14, "color": GRAY},
            ],
            fill=fill,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.25,
            after=4,
        )
    text(
        s,
        rx,
        5.6,
        rw,
        0.8,
        ["AI 모델 = 숫자가 가득 든 큰 파일. 클수록 똑똑하지만 GPU 가 더 든다"],
        size=14,
        color=GRAY,
    )


def stages_slide(d: Deck):
    s = d.content(
        "v3 판독: 13단계 중 AI 는 5곳",
        "나머지 8곳은 프로그램이 읽고 계산한다",
        "v3 가 도면 한 장을 읽는 과정입니다. 13단계로 나뉘어 있고, 주황색 다섯 곳에서만 AI 에게 묻습니다. "
        "표제란 그림을 글자로 받아쓰고, 그 글자를 정리하고, 도면의 뷰와 부품 모양을 판단하고, 전개에 "
        "필요한 치수를 고르고, 3D 후보를 고릅니다. 나머지, 도면 파일을 읽고 계산하는 일은 전부 "
        "프로그램이 합니다. 앞의 DXF 이야기와 같은 분업입니다. 순서는 코드에 고정돼 있어서 01 부터 13 "
        "까지 늘 같은 순서로 돕니다.",
    )
    road(s, 1)
    w, h, gap = 1.62, 1.05, 0.12
    for i, (num, name, model) in enumerate(STAGES):
        row, col = divmod(i, 7)
        card(
            s,
            X0 + col * (w + gap) + (0 if row == 0 else (w + gap) / 2),
            1.95 + row * (h + 0.25),
            w,
            h,
            [
                {"text": num, "size": 12, "bold": True},
                {"text": name, "size": 14, "bold": model},
            ],
            fill=ORANGE if model else LIGHT,
            color=WHITE if model else INK,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            after=0,
            margin=0.04,
        )
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, X0, 4.55, 0.32, 0.32, ORANGE, radius=0.15)
    text(s, X0 + 0.45, 4.5, 3.4, 0.4, ["AI 에게 묻는 곳 (5)"], size=15)
    shape(
        s, MSO_SHAPE.ROUNDED_RECTANGLE, X0 + 4.0, 4.55, 0.32, 0.32, LIGHT, radius=0.15
    )
    text(s, X0 + 4.45, 4.5, 4.5, 0.4, ["프로그램이 하는 곳 (8)"], size=15)
    band(
        s,
        5.35,
        "순서는 고정",
        "01 → 13 을 코드가 늘 같은 순서로 돌린다. AI 는 부품처럼 정해진 곳에서 답만 한다",
        fill=LIGHT,
        color=INK,
        h=1.1,
    )


def synthetic_slide(d: Deck, cone: Path):
    s = d.content(
        "왜 연습용(합성) 도면부터?",
        "로컬 AI 는 Claude 보다 작고, 도면을 보는 눈이 특히 약하다 → 쉬운 시험지부터",
        "그런데 시험은 실제 회사 도면이 아니라 우리가 만든 연습용 도면, 합성 도면으로 했습니다. 이유는 "
        "로컬 AI 가 Claude 보다 작고 약해서입니다. 특히 도면 그림을 보고 알아보는 능력이 더 떨어졌습니다. "
        "실제 회사 도면은 오른쪽처럼 한 장에 부품이 여러 개 있고, 표와 주석이 빽빽합니다. 선과 글자가 "
        "2천에서 6천 개입니다. 연습용 도면은 왼쪽처럼 부품 하나에 뷰 세 개, 선과 글자 66개뿐이고, "
        "정답지도 같이 있습니다. 그래서 쉬운 시험지에서 먼저 되는지 확인하기로 했습니다. 다섯 가지 모양을 "
        "여섯 장씩, 모두 30장입니다.",
    )
    road(s, 1)
    pw, pg = 5.8, 0.533
    for i, (head, fill) in enumerate(
        [("연습용 (합성) 도면", TEAL), ("실제 회사 도면", ORANGE)]
    ):
        chip(
            s,
            X0 + i * (pw + pg),
            1.85,
            pw,
            0.5,
            head,
            fill=fill,
            color=WHITE,
            size=17,
            bold=True,
        )
    pic_w, pic_h = 3.9, 2.7
    picture(s, cone, X0 + (pw - pic_w) / 2, 2.5, pic_w, pic_h)
    real_sheet(s, X0 + pw + pg + (pw - pic_w) / 2, 2.5, pic_w, pic_h)
    facts = [
        [
            ("부품 1개", CHECK),
            ("뷰 3개", CHECK),
            ("선 · 글자 66개", CHECK),
            ("정답지 있음", CHECK),
        ],
        [
            ("부품 여러 개", CROSS),
            ("뷰 후보 최대 13개", CROSS),
            ("선 · 글자 2,000–5,900개", CROSS),
            ("정답지 없음", CROSS),
        ],
    ]
    cw2 = (pw - 0.15) / 2
    for i, rows in enumerate(facts):
        x0 = X0 + i * (pw + pg)
        for j, (label, mark) in enumerate(rows):
            r, c = divmod(j, 2)
            chip(
                s,
                x0 + c * (cw2 + 0.15),
                5.4 + r * 0.6,
                cw2,
                0.5,
                label,
                fill=TEAL_TINT if mark == CHECK else ORANGE_TINT,
                size=14,
                bold=True,
            )


def result_slide(d: Deck):
    s = d.content(
        "결과: 연습 도면 30장 중 27장",
        "도면을 보고 어떤 모양의 부품인지 맞혔다 — 여기까지가 지금",
        "결과입니다. 로컬 AI 가 연습 도면 30장 중 27장에서 부품 모양을 맞혔습니다. 90 % 입니다. 원뿔, "
        "프레스 성형, 절곡판은 여섯 장 모두 맞혔고, 평판은 다섯 장, 원통은 네 장을 맞혔습니다. 동료가 "
        "10월 18일에 돌린 결과입니다. 여기까지가 지금 우리 위치입니다. 쉬운 시험지에서는 로컬 AI 로도 "
        "꽤 된다는 것을 확인했습니다.",
    )
    road(s, 2)
    card(
        s,
        X0,
        1.85,
        3.5,
        4.65,
        [
            {"text": "27", "bold": True, "size": 96, "color": ORANGE},
            {"text": "/ 30 장", "bold": True, "size": 26},
            {"text": "부품 모양 맞힘 · 90 %", "size": 16, "color": GRAY},
        ],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        after=2,
    )
    bar_chart(
        s,
        X0 + 3.8,
        1.85,
        CW - 3.8,
        4.1,
        [name for name, _ in SHAPES],
        [n for _, n in SHAPES],
        [TEAL if n == 6 else ORANGE for _, n in SHAPES],
        number_format='0" / 6"',
        label_size=16,
        gap=45,
    )
    text(
        s,
        X0 + 3.8,
        6.05,
        CW - 3.8,
        0.4,
        ["모양마다 6장 · 로컬 AI (DeepSeek-OCR-2 + Qwen3.6-27B) · 2026-10-18"],
        size=13,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )


def next_overview_slide(d: Deck):
    s = d.content(
        "다음 할 일 두 가지",
        "AI 를 제대로 쓰는 틀을 만들고, 실제 도면으로 간다",
        "그럼 이제 무엇을 할까요. 두 가지입니다. 첫째, 이 수업의 취지대로 AI 를 제대로 써야 합니다. 지금은 "
        "코드가 정해진 곳에서 AI 를 부르는 수준인데, 실제로 AI 를 쓰듯이 AI 가 일하고 그 주변을 받쳐 "
        "주는 틀, 하네스를 만드는 것입니다. 둘째, 연습 도면에서는 되지만 실제 도면에서는 성능이 "
        "떨어집니다. 이걸 끌어올려야 합니다. 그 열쇠가 tool calling 입니다. 그리고 이 둘은 결국 하나로 "
        "이어집니다.",
    )
    road(s, 3)
    items = [
        (
            "1",
            "AI 를 진짜 AI 처럼",
            "하네스 엔지니어링",
            "코드 속 부품 → AI 가 일하고, 하네스가 받친다",
        ),
        (
            "2",
            "실제 도면에서 성능 올리기",
            "tool calling 설계",
            "연습 도면 → 실제 회사 도면",
        ),
    ]
    w = (CW - 0.4) / 2
    for i, (num, head, kind, sub) in enumerate(items):
        x = X0 + i * (w + 0.4)
        card(
            s,
            x,
            1.95,
            w,
            3.45,
            [
                {"text": head, "bold": True, "size": 26},
                {"text": kind, "bold": True, "size": 20, "color": ORANGE},
                {"text": sub, "size": 16, "color": GRAY},
            ],
            margin=0.4,
            after=10,
        ).text_frame.margin_top = Inches(1.35)
        circle(s, x + 0.4, 2.3, 0.85, num, size=26)
    band(s, 5.7, "둘은 이어진다", "좋은 도구 + AI 가 골라 쓰기 + 받쳐 주는 하네스")


def harness_slide(d: Deck):
    s = d.content(
        "할 일 1: 하네스 엔지니어링",
        "AI 를 코드 속 부품으로 부르는 대신, AI 가 일하고 하네스가 받친다",
        "하네스는 원래 말에 채우는 마구를 말합니다. 힘센 말이 원하는 방향으로 일하게 해 주는 장치죠. AI "
        "도 마찬가지입니다. 작은 로컬 AI 라도 일을 잘게 나눠 주고, 답하는 양식을 정해 주고, 답을 검사하고, "
        "모든 걸 기록하고, 점수를 매기면 쓸 만한 결과를 낼 수 있다는 것, 이것이 이 프로젝트가 보이려는 "
        "가설입니다. 아직 증명된 것은 아니고, 연습 도면 27/30 이 첫 근거입니다. 지금 "
        "상태를 표시했습니다. 일 나누기, 답 양식, 기록은 됐습니다. 판독을 13단계로 나눠 단계마다 파일을 "
        "남기고, 단계별로 돌리고 표와 그림으로 보는 도구도 만들었습니다. 평가와 검사는 일부만 있고, AI 가 "
        "순서를 직접 정하는 것은 시제품을 만들었습니다. 다음 장에서 보여 드리겠습니다.",
    )
    road(s, 3)
    card(
        s,
        X0,
        1.85,
        4.0,
        2.0,
        [
            {"text": "하네스 (harness)", "bold": True, "size": 20, "color": TEAL},
            {"text": "말에 채우는 마구.", "size": 16},
            {"text": "원하는 쪽으로 일하게 한다", "size": 16},
        ],
        fill=TEAL_TINT,
        margin=0.3,
        after=4,
    )
    card(
        s,
        X0,
        4.05,
        4.0,
        2.5,
        [
            {"text": "보이려는 것 (가설)", "bold": True, "size": 20, "color": ORANGE},
            {"text": "작은 로컬 AI 라도", "size": 17},
            {"text": "하네스를 잘 만들면", "size": 17},
            {"text": "쓸 만한 결과를 낸다", "size": 17, "bold": True},
        ],
        fill=ORANGE_TINT,
        margin=0.3,
        after=4,
    )
    parts = [
        ("일 나누기", "13단계, 단계마다 파일", CHECK),
        ("답 양식", "정해진 양식으로만 답", CHECK),
        ("기록", "모든 질문과 답을 저장", CHECK),
        ("평가", "연습 30장, 최종 결과만", PART),
        ("검사 · 다시 묻기", "다시 묻기는 아직", PART),
        ("AI 가 순서 정하기", "시제품", PART),
    ]
    bw, bh = 3.0, 1.02
    lx, rx = X0 + 4.35, W - X0 - bw
    dd = 1.5
    cx = (lx + bw + rx) / 2 - dd / 2
    cy = 1.95 + 1.55 + bh / 2 - dd / 2
    for i, (head, sub, mark) in enumerate(parts):
        c, r = divmod(i, 3)
        x = lx if c == 0 else rx
        y = 1.95 + r * 1.55
        connect(s, x + (bw if c == 0 else 0), y + bh / 2, cx + dd / 2, cy + dd / 2)
        done = mark == CHECK
        card(
            s,
            x,
            y,
            bw,
            bh,
            [
                {"text": head, "bold": True, "size": 16},
                {"text": sub, "size": 12, "color": GRAY},
            ],
            fill=TEAL_TINT if done else ORANGE_TINT,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.15,
            after=1,
        ).text_frame.margin_left = Inches(0.65)
        badge(
            s,
            x + 0.12,
            y + bh / 2 - 0.22,
            0.44,
            mark,
            TEAL if done else ORANGE,
            size=15,
        )
    circle(s, cx, cy, dd, "AI", fill=INK, size=28)
    text(
        s,
        lx,
        6.5,
        W - X0 - lx,
        0.35,
        [f"{CHECK} 됨   {PART} 일부 · 시제품"],
        size=12,
        color=GRAY,
        align=PP_ALIGN.CENTER,
    )


def prototype_slide(d: Deck):
    s = d.content(
        "시제품: AI 가 판독 단계를 골라 부른다",
        "명령 한 줄 → 로컬 AI 가 어떤 단계를 돌릴지 정하고 보고했다 (2026-10-23)",
        "하네스의 마지막 조각, AI 가 순서를 정하는 부분의 시제품입니다. v3 의 판독 단계를 AI 가 고를 수 "
        "있는 도구로 만들었습니다. 연습 도면 cone_0019 를 전개도까지 판독해 달라고 한 줄로 명령했더니, "
        "로컬 Qwen 이 먼저 도면을 찾고, 아직 돌린 단계가 없으니 읽기부터 전개까지 돌리고, 결과를 "
        "보고했습니다. AI 가 판단한 건 세 번, 전체 133초였고, 보고한 숫자는 모두 도구 결과에 있던 "
        "값이었습니다. 지어낸 숫자는 없었습니다. 다만 처음부터 끝까지 그대로 돌리는 일에서는 고정 순서와 "
        "차이가 없습니다. 이 방식의 진가는 중간에 막히거나 다시 해야 할 때, 즉 실제 도면에서 나옵니다.",
    )
    road(s, 3)
    card(
        s,
        X0,
        1.9,
        CW,
        0.8,
        [
            [
                {"text": "명령  ", "bold": True, "color": ORANGE},
                {"text": "“cone_0019 도면을 전개도까지 판독하고 결과를 알려줘”"},
            ]
        ],
        fill=ORANGE_TINT,
        size=19,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )
    flow = [
        ("도면 찾기", "list_drawings"),
        ("01 → 10 실행", "run_stages"),
        ("결과 보고", "finish"),
    ]
    w, gap = 3.6, 0.666
    for i, (head, tool) in enumerate(flow):
        x = X0 + i * (w + gap)
        card(
            s,
            x,
            2.95,
            w,
            1.1,
            [
                {"text": head, "bold": True, "size": 18},
                {"text": tool, "size": 13, "color": GRAY, "font": MONO},
            ],
            fill=LIGHT,
            anchor=MSO_ANCHOR.MIDDLE,
            after=2,
        ).text_frame.margin_left = Inches(0.95)
        circle(s, x + 0.2, 3.2, 0.6, str(i + 1), fill=INK, size=16)
        if i < 2:
            arrow(s, x + w + 0.17, 3.33, 0.32, 0.34)
    stats = [("3 번", "AI 판단"), ("133 초", "전체 시간"), ("0 개", "지어낸 숫자")]
    sw = (CW - 2 * 0.3) / 3
    for i, (value, label) in enumerate(stats):
        box = card(
            s,
            X0 + i * (sw + 0.3),
            4.35,
            sw,
            1.25,
            [
                {"text": value, "bold": True, "size": 34, "color": ORANGE},
                {"text": label, "size": 14, "color": GRAY},
            ],
            fill=WHITE,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            after=0,
        )
        box.line.color.rgb = rgb(PALE)
    text(
        s,
        X0,
        5.85,
        CW,
        0.6,
        [
            "AI: Qwen3.6-27B (로컬) · 처음부터 끝까지 그대로 돌리는 일은 고정 순서와 같다 — "
            "진가는 막히거나 다시 해야 할 때"
        ],
        size=14,
        color=GRAY,
    )


def real_slide(d: Deck):
    s = d.content(
        "할 일 2: 실제 도면에서는 막힌다",
        "v3 는 '도면 한 장 = 부품 하나'를 전제로 만들었다. 실제 도면은 그렇지 않다",
        "이제 두 번째 할 일입니다. 같은 v3 에 회사에서 받은 실제 설비 S-2120 의 도면 4장을 넣어 봤습니다. "
        "이 설비의 정답 커팅플랜에는 서로 다른 부품이 22종 있습니다. 그런데 v3 는 도면 한 장에 부품 "
        "하나가 있다고 가정하고 만들어져서, 한 장에 답을 하나만 냅니다. 조립도 한 장에서는 원통 하나를 "
        "읽었는데, 두께를 18 로 읽었습니다. 도면 파일 안에는 't8' 이라고 정확히 적혀 있고 정답도 8 "
        "입니다. 두께가 틀리면 다른 철판을 사게 되고, 펼친 길이도 31 mm 길게 나왔습니다. v2 에서 "
        "말씀드린 첫 단추 문제가 그대로 나타난 것입니다. 나머지 세 장 중 두 장은 한 장에 뷰가 여럿이라 중간 단계에서 멈췄고, 한 장은 "
        "'기타'로 판단해 전개하지 않았습니다. 한 장에 부품 하나라는 전제, 그리고 AI 의 답을 도면 "
        "글자와 대조하지 않는 것, 둘 다 연습 도면에서는 드러나지 않던 문제입니다.",
    )
    road(s, 3)
    big = [
        ("22 종", "정답 부품", ORANGE),
        ("4 장", "입력 도면", INK),
        ("1 개", "v3 가 읽은 부품", GRAY),
    ]
    for i, (value, label, color) in enumerate(big):
        box = card(
            s,
            X0,
            1.85 + i * 1.08,
            4.0,
            0.95,
            [
                [
                    {"text": value, "bold": True, "size": 30, "color": color},
                    {"text": "   " + label, "size": 15, "color": GRAY},
                ]
            ],
            fill=ORANGE_TINT if i == 0 else WHITE,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.25,
        )
        box.line.color.rgb = rgb(PALE)
    text(
        s,
        X0,
        5.1,
        4.0,
        0.4,
        ["설비 S-2120 (회사에서 받은 실제 설비)"],
        size=13,
        color=GRAY,
    )
    rx = X0 + 4.4
    rw = CW - 4.4
    text(
        s,
        rx,
        1.85,
        rw,
        0.5,
        [{"text": "조립도에서 읽은 원통 (SHELL)", "bold": True}],
        size=17,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    rows = [
        ["", "v3 (로컬 AI)", "정답"],
        ["두께", "18 mm", "8 mm"],
        ["펼친 길이", "4,769 mm (+31)", "4,738 mm"],
        ["결과", "두께가 틀리면 철판을 잘못 산다", ""],
    ]

    def style(r, c):
        if r == 0:
            return INK, WHITE, True, 16, PP_ALIGN.CENTER
        bad = c == 1
        return (
            LIGHT if r % 2 else WHITE,
            ORANGE if bad else INK,
            c == 0 or bad,
            20,
            PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER,
        )

    t = table(s, rx, 2.45, rw, 2.4, rows, [2.0, (rw - 2.0) / 2, (rw - 2.0) / 2], style)
    t.cell(3, 1).merge(t.cell(3, 2))
    text(
        s,
        rx,
        4.95,
        rw,
        0.75,
        [
            [
                {"text": "도면 글자는 ", "color": GRAY},
                {"text": "USED TH'K : t8", "bold": True, "font": MONO},
                {"text": "  — 프로그램으로 대조했으면 잡았다", "color": GRAY},
            ],
            {
                "text": "나머지 3장: 2장은 중간 단계에서 멈춤, 1장은 '기타'로 판단",
                "color": GRAY,
            },
        ],
        size=14,
        after=2,
    )
    band(
        s,
        5.75,
        "원인",
        "'한 장 = 부품 하나' 전제, AI 답을 도면 글자와 대조하지 않음 — 연습 도면에선 안 보였다",
    )


def toolcall_slide(d: Deck):
    s = d.content(
        "tool calling: v2 에서 본 그 방식",
        "AI 는 도구 목록을 보고 '이걸 이렇게 불러 줘'라고 요청만 한다. 실행은 프로그램",
        "열쇠는 tool calling 입니다. 사실 이미 보셨습니다. v2 에서 Claude 가 메뉴판을 보고 skill 을 "
        "주문하던 방식, 그게 tool calling 입니다. 프로그램이 도구 목록을 주고, AI 가 하나를 골라 요청하고, "
        "프로그램이 실행해서 결과를 돌려주고, AI 가 다음 도구를 고릅니다. 중요한 점은 AI 는 고르기만 "
        "하고 실행은 늘 프로그램이 한다는 것입니다. 그래서 계산은 언제나 정확합니다. v3 에서는 이걸 "
        "쓰지 않고 순서를 코드에 고정했는데, 이제 로컬 AI 로 tool calling 을 다시 가져옵니다. 시제품에서 "
        "로컬 AI 도 이걸 할 수 있다는 걸 확인했습니다.",
    )
    road(s, 3)
    loop = [
        ("1", "도구 목록을 준다", "이름 · 설명 · 입력 양식", TEAL),
        ("2", "AI 가 고른다", "“이 도구를 이 값으로”", INK),
        ("3", "프로그램이 실행", "계산은 언제나 정확하게", TEAL),
        ("4", "결과를 돌려준다", "→ 다시 2 로", TEAL),
    ]
    bw, bh, bg = 3.3, 1.45, 0.7
    pos = [(X0, 1.9), (X0 + bw + bg, 1.9), (X0 + bw + bg, 4.0), (X0, 4.0)]
    for (num, head, sub, color), (x, y) in zip(loop, pos):
        card(
            s,
            x,
            y,
            bw,
            bh,
            [
                {"text": head, "bold": True, "size": 18},
                {"text": sub, "size": 14, "color": GRAY},
            ],
            fill=LIGHT if color == TEAL else ORANGE_TINT,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.2,
            after=3,
        ).text_frame.margin_left = Inches(0.95)
        circle(s, x + 0.2, y + bh / 2 - 0.3, 0.6, num, fill=color, size=16)
    arrow(s, X0 + bw + 0.17, 2.45, 0.36, 0.36)
    arrow(s, X0 + bw + bg + bw / 2 - 0.18, 3.43, 0.36, 0.5, kind=MSO_SHAPE.DOWN_ARROW)
    arrow(s, X0 + bw + 0.17, 4.55, 0.36, 0.36, kind=MSO_SHAPE.LEFT_ARROW)
    arrow(s, X0 + bw / 2 - 0.18, 3.43, 0.36, 0.5, kind=MSO_SHAPE.UP_ARROW)
    rx = X0 + 2 * bw + bg + 0.35
    rw = CW - (rx - X0)
    history = [
        ("v2", "Claude 가 skill 25개를 tool calling 으로", LIGHT, GRAY),
        ("v3", "순서를 코드에 고정 — tool calling 없음", LIGHT, GRAY),
        ("다음", "로컬 AI 가 tool calling (시제품)", ORANGE_TINT, ORANGE),
    ]
    for i, (name, body, fill, color) in enumerate(history):
        card(
            s,
            rx,
            1.9 + i * 1.22,
            rw,
            1.05,
            [
                {"text": name, "bold": True, "size": 18, "color": color},
                {"text": body, "size": 15},
            ],
            fill=fill,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.2,
            after=2,
        )
    text(
        s,
        rx,
        5.6,
        rw,
        0.8,
        [{"text": "AI 는 실행하지 않는다. 고르기만 한다", "bold": True}],
        size=17,
    )


def toolcall_real_slide(d: Deck):
    s = d.content(
        "tool calling 으로 실제 도면 풀기",
        "작고 확실한 도구를 만들어 두면, AI 가 도면에 맞게 골라 쓰고 결과를 모은다",
        "tool calling 으로 실제 도면을 어떻게 풀 수 있는지 그려 봤습니다. AI 가 도면을 보고 먼저 판단합니다. "
        "'부품이 여러 개 있네, 하나씩 나눠서 보자.' 그리고 부품마다 도구를 골라 씁니다. 부품 영역 자르기, "
        "표제란과 자재표 읽기, 치수 읽기, 모양 판단, 전개 계산, 검사. 청록색은 이미 있는 도구입니다. v3 의 "
        "판독 단계와 v2 의 계산·검사 skill 이죠. 주황색 테두리는 새로 만들 도구입니다. 그리고 검사에서 "
        "이상하면, 예를 들어 AI 가 읽은 두께 18 이 도면 글자 t8 과 다르면, 확대해서 다시 보거나 사람에게 묻는 것도 AI 가 "
        "판단합니다. 마지막에 부품별 결과를 모아 표로 냅니다. 고정 순서로는 이런 대응을 할 수 없습니다.",
    )
    road(s, 3)
    real_sheet(s, X0, 1.95, 2.3, 1.6)
    arrow(s, X0 + 2.45, 2.57, 0.35, 0.36)
    card(
        s,
        X0 + 2.95,
        1.95,
        5.0,
        1.6,
        [
            {"text": "AI 의 판단", "bold": True, "size": 14, "color": ORANGE},
            {
                "text": "“부품이 여러 개다. 하나씩 나눠서 읽자”",
                "size": 18,
                "bold": True,
            },
        ],
        fill=INK,
        color=WHITE,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
        after=4,
    )
    arrow(s, X0 + 8.1, 2.57, 0.35, 0.36)
    tx = X0 + 8.6
    tw = CW - 8.6
    rows = [
        ["부품", "모양", "전개 크기"],
        ["SHELL", "원통", "4738 × 1823"],
        ["TOP HEAD", "헤드", "Ø 1890"],
        ["…", "…", "…"],
    ]

    def style(r, c):
        if r == 0:
            return TEAL, WHITE, True, 12, PP_ALIGN.CENTER
        return LIGHT if r % 2 else WHITE, INK, False, 12, PP_ALIGN.CENTER

    table(s, tx, 1.95, tw, 1.6, rows, [1.25, 0.8, tw - 2.05], style)
    text(s, tx, 3.6, tw, 0.35, ["부품별 결과를 모아 보고"], size=12, color=GRAY)
    card(s, X0, 4.1, CW, 1.65, [""], fill=LIGHT)
    text(
        s,
        X0 + 0.2,
        4.15,
        4.0,
        0.35,
        [{"text": "부품마다 AI 가 고르는 도구", "bold": True}],
        size=13,
        color=GRAY,
    )
    tools = [
        ("부품 영역 자르기", False),
        ("표제란 · 자재표 읽기", True),
        ("치수 읽기", True),
        ("모양 판단", True),
        ("전개 계산", True),
        ("검사", True),
    ]
    n, gap = len(tools), 0.32
    tw2 = (CW - 0.4 - (n - 1) * gap) / n
    for i, (name, have) in enumerate(tools):
        x = X0 + 0.2 + i * (tw2 + gap)
        c = chip(
            s,
            x,
            4.6,
            tw2,
            0.85,
            name,
            fill=TEAL if have else WHITE,
            color=WHITE if have else ORANGE,
            size=14,
            bold=True,
        )
        if not have:
            dashed(c, width=2)
        if i < n - 1:
            arrow(s, x + tw2 + 0.05, 4.87, 0.22, 0.3)
    chip(s, X0, 6.0, 2.0, 0.42, "있는 도구 (v2 · v3)", fill=TEAL, color=WHITE, size=12)
    dashed(
        chip(
            s,
            X0 + 2.15,
            6.0,
            1.8,
            0.42,
            "새로 만들 도구",
            fill=WHITE,
            color=ORANGE,
            size=12,
        )
    )
    text(
        s,
        X0 + 4.2,
        5.97,
        CW - 4.2,
        0.5,
        [
            [
                {"text": "이상하면  ", "bold": True, "color": ORANGE},
                {
                    "text": "확대해서 다시 보기 · 다른 도구로 · 사람에게 묻기 — 이것도 AI 가 고른다"
                },
            ]
        ],
        size=15,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def outlook_slide(d: Deck):
    s = d.content(
        "전망: 함수로 만들 수 있으면, AI 가 쓴다",
        "필요한 기능을 파이썬 함수로 만들고 AI 가 적재적소에 부르면 → 최종 목표까지",
        "마지막으로 전망입니다. 우리에게 필요한 기능을 파이썬 함수로 만들 수만 있다면, 그건 AI 가 쓸 수 "
        "있는 도구가 됩니다. AI 가 그 도구를 적재적소에 부르고 결과를 모으고, 하네스가 검사하고 기록하고 "
        "점수를 매기면, 도면에서 커팅플랜과 구매의뢰까지 갈 수 있다고 봅니다. 그리고 재료는 이미 꽤 "
        "있습니다. 도면 읽기 도구는 v3 의 13단계가 있고, 전개·분할·네스팅·구매의뢰를 계산하는 도구는 v2 "
        "에서 만들어 둔 것이 사람이 넣은 부품으로 끝까지 돌아갑니다. 로컬 AI 가 도구를 골라 쓰는 것도 "
        "시제품으로 확인했습니다. 남은 일은 실제 도면에 필요한 도구를 더하고, 단계마다 점수를 매기는 "
        "것입니다.",
    )
    road(s, 3)
    eq = [
        ("파이썬 함수", "필요한 기능 하나씩", LIGHT, INK),
        ("AI 가 골라 부른다", "tool calling", ORANGE_TINT, ORANGE),
        ("하네스", "검사 · 기록 · 평가", TEAL_TINT, TEAL),
        ("도면 → 구매의뢰", "최종 목표", INK, WHITE),
    ]
    w, gap = 2.55, 0.64
    for i, (head, sub, fill, color) in enumerate(eq):
        x = X0 + i * (w + gap)
        card(
            s,
            x,
            1.95,
            w,
            1.5,
            [
                {"text": head, "bold": True, "size": 19, "color": color},
                {"text": sub, "size": 14, "color": "B8C4C9" if fill == INK else GRAY},
            ],
            fill=fill,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            after=3,
        )
        if i < 3:
            text(
                s,
                x + w,
                1.95,
                gap,
                1.5,
                ["=" if i == 2 else "+"],
                size=36,
                bold=True,
                color=GRAY,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
            )
    have = [
        ("도면 읽기 도구", "v3 판독 13단계"),
        ("계산 도구", "v2 전개 · 분할 · 네스팅 · 구매의뢰 — 끝까지 돈다"),
        ("AI 가 고르기", "로컬 AI 시제품"),
    ]
    todo = [
        ("실제 도면용 도구", "부품 영역 자르기, 다시 보기 …"),
        ("단계별 점수", "어디서 틀리는지 숫자로"),
    ]
    hw = 7.2
    text(
        s,
        X0,
        3.75,
        hw,
        0.4,
        [{"text": "이미 있는 것", "bold": True}],
        size=15,
        color=TEAL,
    )
    for i, (head, sub) in enumerate(have):
        y = 4.2 + i * 0.78
        badge(s, X0, y + 0.05, 0.52, CHECK, TEAL, size=18)
        text(
            s,
            X0 + 0.7,
            y,
            hw - 0.7,
            0.65,
            [[{"text": head + "  ", "bold": True}, {"text": sub, "color": GRAY}]],
            size=16,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    tx = X0 + hw + 0.4
    tw = CW - hw - 0.4
    text(
        s, tx, 3.75, tw, 0.4, [{"text": "남은 것", "bold": True}], size=15, color=ORANGE
    )
    for i, (head, sub) in enumerate(todo):
        card(
            s,
            tx,
            4.2 + i * 1.17,
            tw,
            1.0,
            [
                {"text": head, "bold": True, "size": 17},
                {"text": sub, "size": 14, "color": GRAY},
            ],
            fill=ORANGE_TINT,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.25,
            after=2,
        )


def closing_slide(d: Deck):
    s = d.slide(
        "정리하겠습니다. v2 에서는 Claude 에게 도면부터 구매의뢰까지 전부 맡겼는데, 일이 너무 커서 잘 "
        "안 됐습니다. v3 에서는 우리 서버의 로컬 AI 로 도면 읽기에 집중했고, 연습 도면 30장 중 27장을 "
        "맞혔습니다. 여기가 지금입니다. 다음은 하네스와 tool calling 으로 실제 도면을 읽고, 이미 있는 "
        "계산 도구까지 이어서 끝까지 가는 것입니다. 질문 받겠습니다.",
        dark=True,
    )
    text(s, X0, 0.8, CW, 0.8, ["정리"], size=36, bold=True, color=WHITE)
    points = [
        ("v2", "Claude 에게 전부 맡겼다 → 일이 너무 컸다"),
        ("v3", "로컬 AI, 도면 읽기에 집중 → 연습 도면 27 / 30"),
        ("지금", "여기. 실제 도면에서는 막힌다"),
        ("다음", "하네스 + tool calling 으로 실제 도면, 끝까지"),
    ]
    for i, (label, point) in enumerate(points):
        y = 1.9 + i * 1.05
        chip(
            s,
            X0,
            y + 0.05,
            1.3,
            0.6,
            label,
            fill=ORANGE if label == "지금" else "34444C",
            color=WHITE,
            size=18,
            bold=True,
        )
        text(
            s,
            X0 + 1.6,
            y,
            CW - 1.6,
            0.7,
            [point],
            size=22,
            color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    text(s, X0, 6.3, CW, 0.6, ["질문"], size=26, bold=True, color=ORANGE)


def main() -> None:
    cone = render_cone()
    d = Deck(OUT, TITLE, FOOTER)
    title_slide(d)
    map_slide(d)
    v2_slide(d)
    skill_slide(d)
    claude_loop_slide(d)
    v2_result_slide(d)
    why_move_slide(d)
    v3_shift_slide(d)
    server_slide(d)
    stages_slide(d)
    synthetic_slide(d, cone)
    result_slide(d)
    next_overview_slide(d)
    harness_slide(d)
    prototype_slide(d)
    real_slide(d)
    toolcall_slide(d)
    toolcall_real_slide(d)
    outlook_slide(d)
    closing_slide(d)
    print(f"{d.count} slides -> {d.save()}")


if __name__ == "__main__":
    main()
