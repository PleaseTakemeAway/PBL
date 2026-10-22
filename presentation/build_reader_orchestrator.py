"""Build presentation/reader_orchestrator.pptx: the reader orchestrator part of the talk.

PYTHONPATH=. uv run --with python-pptx python presentation/build_reader_orchestrator.py

The slides are written here, not edited in PowerPoint, so a change shows up as a text diff
and branches can be merged; rebuild the .pptx after a merge. Speaker notes are in
slide.notes. Every number on the slides comes from docs/reader/orchestrator.md (reader-agent
branch) and the 2026-10-23 cone_0019 run on mlic.
"""

import datetime
import io
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("reader_orchestrator.pptx")
FIXED_TIME = datetime.datetime(2026, 1, 1)
FOOTER = "PBL cutplan_agent · 도면 판독 orchestrator · 2026-10-24"

FONT = "Malgun Gothic"  # 맑은 고딕: Korean font that ships with Office on Windows
MONO = "Consolas"

INK = "1F2A30"  # dark slate: title/closing background, body text
TEAL = "2E6F73"  # steel teal: structure, "done"
ORANGE = "E8743B"  # cutting-torch orange: model stages, highlights
GRAY = "5F6B73"
LIGHT = "EEF3F4"  # card tint
TEAL_TINT = "DCEBEC"
ORANGE_TINT = "FBE7DC"
WHITE = "FFFFFF"
CODE_BG = "263238"
CODE_FG = "E6EDF0"

W, H = 13.333, 7.5
X0, CW = 0.6, 12.133  # content left edge and width


# ---------------------------------------------------------------------- primitives
def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def style_run(run, size, color, bold=False, font=FONT, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):  # Korean glyphs use the East Asian font, not latin
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", font)


def set_bullet(paragraph, color):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", str(Inches(0.28)))
    ppr.set("indent", str(-Inches(0.22)))
    clr = ppr.makeelement(qn("a:buClr"), {})
    srgb = clr.makeelement(qn("a:srgbClr"), {"val": color})
    clr.append(srgb)
    ppr.append(clr)
    ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "•"}))


def write(
    tf,
    paras,
    *,
    size=16,
    color=INK,
    bold=False,
    font=FONT,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    after=6,
    margin=0.0,
):
    """paras: list of paragraphs; a paragraph is a str, a run dict, or a list of runs.
    A dict paragraph may set bullet=True and its own size/color/bold/font/italic."""
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        runs = para if isinstance(para, list) else [para]
        bullet = False
        for r in runs:
            r = {"text": r} if isinstance(r, str) else r
            bullet = bullet or r.get("bullet", False)
            run = p.add_run()
            run.text = r["text"]
            style_run(
                run,
                r.get("size", size),
                r.get("color", color),
                r.get("bold", bold),
                r.get("font", font),
                r.get("italic", False),
            )
        if bullet:
            set_bullet(p, ORANGE)
    return tf


def text(slide, x, y, w, h, paras, **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    write(box.text_frame, paras, **kw)
    return box


def shape(slide, kind, x, y, w, h, fill, *, line=None, radius=0.06):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def card(slide, x, y, w, h, paras, *, fill=LIGHT, margin=0.18, **kw):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill)
    write(s.text_frame, paras, margin=margin, **kw)
    return s


def circle(slide, x, y, d, label, fill=ORANGE, color=WHITE, size=16):
    s = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)
    write(
        s.text_frame,
        [{"text": label, "bold": True}],
        size=size,
        color=color,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    return s


def arrow(slide, x, y, w=0.3, h=0.32, fill=ORANGE, kind=MSO_SHAPE.RIGHT_ARROW):
    return shape(slide, kind, x, y, w, h, fill)


def code(slide, x, y, w, h, lines, size=11):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, CODE_BG, radius=0.03)
    write(
        s.text_frame,
        [{"text": line} for line in lines],
        size=size,
        color=CODE_FG,
        font=MONO,
        after=2,
        margin=0.2,
    )
    return s


def chip(slide, x, y, w, h, label, *, fill=LIGHT, color=INK, size=13, bold=False):
    s = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, radius=0.15)
    write(
        s.text_frame,
        [{"text": label, "bold": bold}],
        size=size,
        color=color,
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        after=0,
        margin=0.05,
    )
    return s


# ---------------------------------------------------------------------- slide frames
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.count = 0

    def slide(self, notes, *, dark=False):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(INK if dark else WHITE)
        s.notes_slide.notes_text_frame.text = notes
        self.count += 1
        return s

    def content(self, title, message, notes):
        """A light slide with the title, a one-line message under it and the footer."""
        s = self.slide(notes)
        text(s, X0, 0.45, CW, 0.75, [title], size=34, bold=True, anchor=MSO_ANCHOR.TOP)
        text(s, X0, 1.15, CW, 0.45, [message], size=17, color=GRAY)
        text(s, X0, 6.95, 9.0, 0.3, [FOOTER], size=10, color=GRAY)
        text(
            s,
            W - X0 - 1.0,
            6.95,
            1.0,
            0.3,
            [str(self.count)],
            size=10,
            color=GRAY,
            align=PP_ALIGN.RIGHT,
        )
        return s

    def save(self):
        props = self.prs.core_properties
        props.title = "LLM이 도면 판독 파이프라인을 직접 운영하게 하기"
        props.author = "PBL cutplan_agent"
        props.created = props.modified = FIXED_TIME
        props.revision = 1
        buffer = io.BytesIO()
        self.prs.save(buffer)
        reproducible(buffer.getvalue(), OUT)
        return OUT


def reproducible(data: bytes, path: Path) -> None:
    """Rewrite the zip with fixed entry times, so the same slides give the same bytes and
    a rebuild without changes leaves git clean."""
    with zipfile.ZipFile(io.BytesIO(data)) as src:
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                entry = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
                entry.compress_type = zipfile.ZIP_DEFLATED
                dst.writestr(entry, src.read(info.filename))


# ---------------------------------------------------------------------- slides
STAGES = [
    ("01", "read", False),
    ("02", "title_regex", False),
    ("03", "title_ocr", True),
    ("04", "title_fields", True),
    ("05", "views", False),
    ("06", "view_labels", True),
    ("07", "sheet", False),
    ("08", "document", False),
    ("09", "unfold_params", True),
    ("10", "develop", False),
    ("11", "reconstruct", False),
    ("12", "confirm", True),
    ("13", "export", False),
]


def title_slide(d: Deck):
    s = d.slide(
        "이번 파트에서는 도면 판독기(reader)에 LLM orchestrator 를 붙인 작업을 소개합니다. "
        "판독 13단계를 모델이 tool 로 골라 부르게 만들었고, 로컬 모델 Qwen3.6-27B 로 실제 "
        "도면 한 장을 끝까지 돌려 봤습니다. 순서는 왜 필요한지, 무엇을 만들었는지, 돌려 보니 "
        "어땠는지, 다음에 무엇을 할지입니다.",
        dark=True,
    )
    text(
        s,
        X0,
        2.2,
        CW,
        1.6,
        ["LLM이 도면 판독 파이프라인을", "직접 운영하게 하기"],
        size=40,
        bold=True,
        color=WHITE,
        after=4,
    )
    text(
        s,
        X0,
        4.0,
        CW,
        0.5,
        ["Tool calling 기반 판독 orchestrator — 로컬 LLM 하네스 엔지니어링"],
        size=20,
        color=ORANGE,
    )
    text(
        s,
        X0,
        6.4,
        CW,
        0.4,
        ["PBL cutplan_agent · 2026-10-24"],
        size=14,
        color="B8C4C9",
    )


def goal_slide(d: Deck):
    s = d.content(
        "목표: 값싼 로컬 LLM + 좋은 하네스",
        "모델을 키우는 대신, 모델 바깥의 장치로 결과를 낸다",
        "프로젝트의 목표를 다시 짚겠습니다. 비싼 모델 하나에 맡기는 대신, 값싼 로컬 모델 여러 개와 "
        "잘 만든 하네스로 원하는 결과를 내는 것입니다. 하네스는 모델 바깥에서 모델이 일을 잘하게 "
        "만드는 장치 전부를 말합니다. 그중 작업 분해와 검증은 판독기에 이미 꽤 있는데, "
        "orchestration, 즉 어떤 단계를 언제 부를지를 로컬 모델이 정하는 부분은 비어 있었습니다. "
        "오늘 이야기는 그 빈칸입니다.",
    )
    items = [
        (
            "작업 분해",
            "큰 일을 단계로 나누고, 모델에게는 좁고 확인 가능한 판단만 맡긴다",
            "있음 · 13단계",
            TEAL,
        ),
        (
            "검증과 폴백",
            "모델 답을 코드로 검사하고, 틀리면 결정적 방법이나 사람 확인으로",
            "있음 · review",
            TEAL,
        ),
        (
            "orchestration",
            "어떤 단계를 어떤 모델로 언제 다시 부를지 정한다",
            "없음 → 오늘 주제",
            ORANGE,
        ),
    ]
    for i, (head, body, state, color) in enumerate(items):
        x = X0 + i * (3.9 + 0.2)
        card(
            s,
            x,
            1.95,
            3.9,
            3.0,
            [
                {"text": head, "bold": True, "size": 22, "color": color},
                {"text": body, "size": 16},
            ],
            margin=0.3,
            after=10,
        )
        chip(s, x + 0.3, 4.3, 2.4, 0.42, state, fill=color, color=WHITE, size=14)
    card(
        s,
        X0,
        5.3,
        CW,
        1.2,
        [
            [
                {"text": "성공 기준  ", "bold": True, "color": ORANGE},
                {
                    "text": "orchestration 도 로컬 모델이 하고, 같은 하네스에서 Claude 와 "
                    "비슷한 결과를 숫자로 보인다 (docs/harness_engineering.md)"
                },
            ]
        ],
        fill=ORANGE_TINT,
        size=16,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )


def reader_slide(d: Deck):
    s = d.content(
        "지금 판독기: LLM은 부품이다",
        "13단계가 코드에 고정된 순서로 돌고, LLM 은 5곳에서 한 번씩 묻고 답을 받는다",
        "지금 판독기 구조입니다. DXF 도면 한 장을 13단계로 처리합니다. experiments.sh 는 vLLM 서버를 "
        "띄우고 main.py --batch 를 부르고, 그 안의 run_drawing 이 13단계를 정해진 순서대로 돕니다. "
        "주황색 5개 단계만 모델을 씁니다. 표제란 OCR 하나와 추론 넷인데, 각각 한 번 묻고 JSON 답을 "
        "받을 뿐입니다. 무엇을 할지 고르거나 도구를 부르지 않습니다. 그래서 지금 판독기는 agent 가 "
        "아니라 LLM 을 부품으로 쓰는 고정 workflow 입니다.",
    )
    w, h, gap = 1.62, 0.95, 0.12
    for i, (num, name, model) in enumerate(STAGES):
        row, col = divmod(i, 7)
        x = X0 + col * (w + gap) + (0 if row == 0 else (w + gap) / 2)
        y = 1.95 + row * (h + 0.2)
        shp = shape(
            s,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x,
            y,
            w,
            h,
            ORANGE if model else LIGHT,
            radius=0.12,
        )
        write(
            shp.text_frame,
            [
                {"text": num, "size": 12, "bold": True},
                {"text": name, "size": 13},
            ],
            color=WHITE if model else INK,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            after=0,
            margin=0.04,
        )
    chip(s, X0, 4.25, 0.32, 0.32, "", fill=ORANGE)
    text(
        s,
        X0 + 0.45,
        4.22,
        6.0,
        0.4,
        ["모델 사용 단계 — OCR 1곳(03), 추론 4곳(04 · 06 · 09 · 12)"],
        size=14,
        color=GRAY,
    )
    code(
        s,
        X0,
        4.85,
        5.9,
        1.0,
        [
            "experiments.sh → main.py --batch",
            "  → run_drawing(): 01 → 02 → … → 13 (순서 고정)",
        ],
        size=13,
    )
    card(
        s,
        X0 + 6.2,
        4.85,
        CW - 6.2,
        1.0,
        [
            [
                {"text": "빈칸  ", "bold": True, "color": ORANGE},
                {"text": "순서를 정하는 LLM 이 없다 (로드맵 6번)"},
            ]
        ],
        fill=ORANGE_TINT,
        size=17,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )
    text(
        s,
        X0,
        6.05,
        CW,
        0.5,
        [
            "v2 워크오더 에이전트에는 Claude orchestrator 가 있었지만 Claude 전용 기능"
            "(tool_runner)에 묶여 있다"
        ],
        size=14,
        color=GRAY,
    )


def toolcall_slide(d: Deck):
    s = d.content(
        "Tool calling: 모델은 고르기만 한다",
        "모델은 코드를 실행하지 않는다. 다음에 부를 함수와 인자를 JSON 으로 요청할 뿐이다",
        "tool calling 이 무엇인지 짧게 짚겠습니다. 프로그램이 모델에게 질문과 함께 도구 목록을 "
        "보냅니다. 도구마다 이름, 설명, 인자 형식이 JSON Schema 로 들어갑니다. 모델은 답 대신 "
        "'run_views 를 이 인자로 불러 달라'는 JSON 을 돌려주고, 실제 실행은 우리 프로그램이 합니다. "
        "결과를 다시 모델에게 주면 모델이 다음 도구를 고르고, 더 부를 도구가 없으면 최종 답을 "
        "씁니다. 핵심은 모델이 하는 일이 '고르기'뿐이라는 점입니다. 계산과 판독은 여전히 코드가 "
        "합니다.",
    )
    steps = [
        ("1", "질문 + 도구 목록", "이름 · 설명 · 인자 JSON Schema 를 같이 보낸다"),
        ("2", "모델이 JSON 요청", "이 함수를 이 인자로 불러 달라"),
        ("3", "프로그램이 실행", "단계 함수가 돌고 파일을 남긴다"),
        ("4", "결과를 돌려준다", "모델이 다음 도구를 고른다. 없으면 최종 답"),
    ]
    w, gap = 2.75, 0.38
    for i, (num, head, body) in enumerate(steps):
        x = X0 + i * (w + gap)
        card(
            s,
            x,
            1.85,
            w,
            2.25,
            [
                {"text": head, "bold": True, "size": 18},
                {"text": body, "size": 15, "color": GRAY},
            ],
            margin=0.22,
            after=8,
        ).text_frame.margin_top = Inches(0.75)
        circle(s, x + 0.22, 2.0, 0.5, num)
        if i < 3:
            arrow(s, x + w + 0.05, 2.8, 0.28, 0.32)
    text(s, X0, 4.25, 5.9, 0.35, ["프로그램이 보내는 도구 정의"], size=13, color=GRAY)
    code(
        s,
        X0,
        4.6,
        5.9,
        2.1,
        [
            '{"name": "run_view_labels",',
            ' "description": "Stage 06 view_labels: name every',
            '   view cluster. Model: reasoning ...",',
            ' "parameters": {"type": "object", "properties": {',
            '   "work": {"type": "string"},',
            '   "projection_default": {"enum": ["first", "third"]}}}}',
        ],
    )
    text(s, X0 + 6.2, 4.25, 5.9, 0.35, ["모델이 돌려주는 요청"], size=13, color=GRAY)
    code(
        s,
        X0 + 6.2,
        4.6,
        CW - 6.2,
        2.1,
        [
            '{"tool_calls": [{"function": {',
            '   "name": "run_view_labels",',
            '   "arguments": {"work": "cone_0019",',
            '                 "projection_default": "third"}}}]}',
        ],
    )


def mcp_slide(d: Deck):
    s = d.content(
        "Tool calling, MCP, Agent Skills",
        "셋은 다른 층의 이야기다. 로컬 orchestrator 에 필요한 것은 tool calling 이다",
        "자주 섞이는 세 용어를 정리했습니다. tool calling 은 모델 API 의 기능으로, 모델이 함수 호출을 "
        "요청하는 방식입니다. MCP 는 도구를 별도 서버로 떼어 내서 여러 프로그램이 같은 규격으로 쓰게 "
        "하는 프로토콜입니다. MCP 를 써도 모델에게는 결국 tool calling 형식으로 전달됩니다. Agent "
        "Skills 는 SKILL.md 라는 지침 문서 묶음을 모델이 골라 읽는 규격입니다. 우리 레포의 skills "
        "폴더와는 이름만 같습니다. 우리 레포는 코드가 단계마다 붙일 문서를 고정해 둡니다. 판정은 "
        "tool calling 으로 충분하다는 것입니다. 도구를 쓰는 프로그램이 orchestrator 하나뿐이고 도구가 "
        "같은 프로세스의 파이썬 함수이기 때문입니다.",
    )
    cols = [
        (
            "Tool calling",
            TEAL,
            [
                ("무엇", "모델이 '이 함수를 이 인자로' JSON 요청"),
                ("어디서", "우리 프로그램 안에서 함수 직접 호출"),
                ("우리 프로젝트", "로컬 orchestrator 의 핵심 → 이번에 구현"),
            ],
        ),
        (
            "MCP",
            GRAY,
            [
                ("무엇", "도구를 별도 서버로 떼어 여러 클라이언트가 공유 (JSON-RPC)"),
                ("어디서", "별도 서버 프로세스, host 가 tool calling 으로 변환"),
                ("우리 프로젝트", "도구를 쓰는 프로그램이 하나 → 지금은 불필요"),
            ],
        ),
        (
            "Agent Skills",
            GRAY,
            [
                ("무엇", "SKILL.md 지침 묶음, 모델이 필요할 때 골라 읽음"),
                ("어디서", "에이전트가 문서를 읽는다"),
                ("우리 프로젝트", "skills/ 는 코드가 단계별로 골라 넣는다 (다른 것)"),
            ],
        ),
    ]
    w, gap = 3.9, 0.215
    for i, (head, color, rows) in enumerate(cols):
        x = X0 + i * (w + gap)
        paras = [{"text": head, "bold": True, "size": 22, "color": color}]
        for label, body in rows:
            paras.append({"text": label, "bold": True, "size": 13, "color": GRAY})
            paras.append({"text": body, "size": 15})
        card(
            s,
            x,
            1.85,
            w,
            3.55,
            paras,
            fill=TEAL_TINT if i == 0 else LIGHT,
            margin=0.25,
            after=3,
        )
    card(
        s,
        X0,
        5.6,
        CW,
        1.05,
        [
            [
                {"text": "판정  ", "bold": True, "color": ORANGE},
                {
                    "text": "tool calling 으로 충분하다. MCP 는 Claude Code 같은 다른 클라이언트도 "
                    "같은 단계를 써야 할 때 — 이미 만든 tool 목록을 감싸면 된다"
                },
            ]
        ],
        fill=ORANGE_TINT,
        size=16,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )


def architecture_slide(d: Deck):
    s = d.content(
        "만든 것: 판독 orchestrator",
        "단계 함수는 그대로 두고, 그 위에 '순서를 정하는 모델'만 얹었다",
        "만든 것의 구조입니다. 왼쪽에서 자연어 명령이 들어오면 orchestrator 루프가 모델에게 묻고, "
        "모델이 고른 tool 을 실행합니다. tool 은 17개이고, 그중 13개가 판독 단계 하나씩입니다. "
        "tool 은 기존 단계 함수를 그대로 부르기 때문에, 고정 순서로 돌린 것과 같은 파일이 작업 "
        "폴더에 남습니다. 달라진 것은 순서를 누가 정하느냐 하나뿐입니다. 이게 중요한 이유는 "
        "나중에 고정 순서와 orchestrator 를 비교할 때 차이가 순서 결정에서만 나오기 때문입니다. "
        "모델 칸에는 로컬 Qwen 이나 Claude 를 끼울 수 있고, 모든 턴은 trajectory 파일에 남습니다.",
    )
    y, h = 3.05, 1.3
    gap = (CW - (1.9 + 2.25 + 2.0 + 2.25 + 2.1)) / 4
    boxes = [
        ("자연어 명령", '"cone_0019 판독해줘"', ORANGE_TINT, 1.9),
        ("orchestrator", "루프 · orchestrator.py", TEAL_TINT, 2.25),
        ("tool 17개", "tools.py", LIGHT, 2.0),
        ("판독 단계 13개", "stages/*.py 의 run()", LIGHT, 2.25),
        ("작업 폴더", "stages/NN_*.json", LIGHT, 2.1),
    ]
    x = X0
    for i, (head, sub, fill, w) in enumerate(boxes):
        card(
            s,
            x,
            y,
            w,
            h,
            [
                {"text": head, "bold": True, "size": 17},
                {"text": sub, "size": 13, "color": GRAY},
            ],
            fill=fill,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.15,
            after=4,
        )
        if i < 4:
            kind = MSO_SHAPE.LEFT_RIGHT_ARROW if i == 1 else MSO_SHAPE.RIGHT_ARROW
            arrow(s, x + w + 0.03, y + h / 2 - 0.15, 0.35, 0.3, kind=kind)
        if i == 1:
            loop_x, loop_w = x, w
        if i == 3:
            stage_x = x
        x += w + gap
    card(
        s,
        loop_x,
        1.75,
        loop_w,
        0.8,
        [
            {"text": "모델", "bold": True, "size": 15},
            {"text": "Qwen3.6-27B 또는 Claude", "size": 13, "color": GRAY},
        ],
        fill=ORANGE_TINT,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.12,
        after=2,
    )
    arrow(s, loop_x + loop_w / 2 - 0.15, 2.58, 0.3, 0.44, kind=MSO_SHAPE.UP_DOWN_ARROW)
    arrow(s, loop_x + loop_w / 2 - 0.15, 4.38, 0.3, 0.3, kind=MSO_SHAPE.DOWN_ARROW)
    card(
        s,
        loop_x,
        4.72,
        loop_w,
        0.85,
        [
            {"text": "기록", "bold": True, "size": 15},
            {"text": "trajectory.jsonl", "size": 13, "color": GRAY},
        ],
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.12,
        after=2,
    )
    card(
        s,
        stage_x,
        4.72,
        2.25,
        0.85,
        [
            {"text": "단계 안의 모델", "bold": True, "size": 15},
            {"text": "OCR · 추론, 기존 그대로", "size": 13, "color": GRAY},
        ],
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.12,
        after=2,
    )
    text(
        s,
        X0,
        5.95,
        CW,
        0.6,
        [
            [
                {"text": "같은 단계 함수 → 같은 결과 파일.  ", "bold": True},
                {
                    "text": "고정 순서(run.py)와 다른 것은 '순서를 누가 정하나' 하나뿐이라 둘을 "
                    "공정하게 비교할 수 있다",
                    "color": GRAY,
                },
            ]
        ],
        size=15,
    )


def tools_slide(d: Deck):
    s = d.content(
        "Tool 17개를 만든 원칙",
        "약한 모델도 실수하기 어렵게, 실수해도 루프가 멈추지 않게",
        "tool 을 만들 때 지킨 원칙 네 가지입니다. 첫째, 결과를 작게 돌려줍니다. 첫 단계 파일에는 DXF "
        "엔티티가 전부 들어 있어서 그대로 주면 로컬 모델의 16k 컨텍스트가 바로 찹니다. 그래서 상태와 "
        "요약, 모델 호출 성공 여부만 줍니다. 둘째, 오류도 결과로 돌려줍니다. 없는 tool 이나 틀린 "
        "인자를 줘도 예외로 죽지 않고 모델이 오류를 보고 고칠 수 있습니다. 셋째, 단계 안에서 쓰는 OCR, "
        "추론 모델은 시작할 때 고정하고 orchestrator 는 순서만 고릅니다. 넷째, 작업 폴더는 출력 폴더 "
        "아래로만 제한합니다. 오른쪽이 tool 목록입니다.",
    )
    rules = [
        (
            "결과는 작게",
            "단계 파일 전체 대신 상태 · 요약 · 모델 호출 결과만. "
            "01_read.json 은 DXF 엔티티 전부 → 결과는 1,000자 미만",
        ),
        (
            "오류도 결과로",
            '없는 tool, 틀린 인자, 앞 단계 없음(not_ready), 단계 오류 → {"error": …}. '
            "루프는 멈추지 않는다",
        ),
        (
            "모델은 순서만",
            "단계 안의 OCR · 추론 모델은 시작할 때 고정. orchestrator 는 단계와 옵션만 고른다",
        ),
        (
            "작업 폴더 밖 금지",
            "work 는 --output-dir 아래만. 'cone_0019' 처럼 짧게 써도 된다",
        ),
    ]
    w, h = 4.0, 2.2
    for i, (head, body) in enumerate(rules):
        row, col = divmod(i, 2)
        x = X0 + col * (w + 0.25)
        y = 1.85 + row * (h + 0.25)
        card(
            s,
            x,
            y,
            w,
            h,
            [
                {"text": head, "bold": True, "size": 19, "color": TEAL},
                {"text": body, "size": 15},
            ],
            margin=0.25,
            after=8,
        )
    lx = X0 + 2 * w + 0.5
    lw = CW - (lx - X0)
    card(
        s,
        lx,
        1.85,
        lw,
        4.65,
        [{"text": "tool 목록", "bold": True, "size": 19, "color": TEAL}],
        fill=TEAL_TINT,
        margin=0.25,
    )
    names = [
        ("list_drawings", "도면과 작업 폴더"),
        ("stage_status", "13단계 상태"),
        ("show_stage", "요약 · review · 모델 원문"),
        ("run_<단계> × 13", "단계 하나"),
        ("run_stages", "여러 단계를 차례로"),
    ]
    for i, (name, what) in enumerate(names):
        y = 2.6 + i * 0.75
        chip(
            s,
            lx + 0.25,
            y,
            lw - 0.5,
            0.42,
            name,
            fill=ORANGE if "×" in name else WHITE,
            color=WHITE if "×" in name else INK,
            size=14,
            bold=True,
        )
        text(s, lx + 0.3, y + 0.43, lw - 0.6, 0.3, [what], size=12, color=GRAY)


def modes_slide(d: Deck):
    s = d.content(
        "Tool을 부르는 세 가지 방식",
        "같은 루프에 모델 호출 방식만 바꿔 끼운다. 로컬과 Claude 를 같은 조건에서 비교하기 위해서다",
        "모델이 tool 을 부르는 방식은 세 가지로 만들었습니다. native 는 OpenAI 호환 표준 tool calling "
        "입니다. 그런데 지금 띄워 둔 vLLM 서버는 tool parser 없이 떠 있어서 400 을 돌려줍니다. "
        "그래서 schema 모드를 만들었습니다. tool 목록을 JSON schema 의 anyOf 로 바꿔서, 모델의 응답이 "
        "반드시 tool 호출 하나의 형식이 되게 디코딩 단계에서 강제합니다. 약한 모델이 형식을 틀릴 수가 "
        "없습니다. 세 번째는 Claude 입니다. Claude 전용 tool_runner 를 쓰지 않고 같은 루프를 쓴 이유는 "
        "로컬과 Claude 를 같은 조건에서 비교하기 위해서입니다.",
    )
    rows = [
        ["방식", "모델이 tool 을 부르는 법", "서버 조건", "상태"],
        [
            "native",
            "OpenAI 호환 tools → tool_calls",
            "vLLM --enable-auto-tool-choice\n--tool-call-parser",
            "서버 재기동 필요\n(지금은 400)",
        ],
        [
            "schema",
            "응답마다 {thought, tool, arguments}\nJSON 하나를 schema 로 강제",
            "structured outputs\n(지금 서버 그대로)",
            "동작 확인",
        ],
        [
            "anthropic",
            "Claude tool use, 같은 루프",
            "ANTHROPIC_API_KEY",
            "미실행 (비용)",
        ],
    ]
    widths = [1.7, 4.5, 3.6, 2.333]
    table = s.shapes.add_table(
        len(rows), 4, Inches(X0), Inches(1.85), Inches(CW), Inches(2.6)
    ).table
    for c, w in enumerate(widths):
        table.columns[c].width = Inches(w)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(INK)
            elif r == 2:
                cell.fill.fore_color.rgb = rgb(TEAL_TINT)
            else:
                cell.fill.fore_color.rgb = rgb(LIGHT if r % 2 else WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            write(
                cell.text_frame,
                [{"text": line} for line in value.split("\n")],
                size=14 if r else 14,
                bold=r == 0 or c == 0,
                color=WHITE if r == 0 else (TEAL if c == 3 and r == 2 else INK),
                after=0,
                margin=0.1,
            )
    code(
        s,
        X0,
        4.75,
        6.3,
        1.85,
        [
            "// schema 모드의 응답 형식 (vLLM xgrammar 가 강제)",
            '{"anyOf": [',
            '  {"tool": {"const": "run_views"}, "arguments": {…}},',
            "  … tool 17개 …,",
            '  {"tool": {"const": "finish"}, "arguments": {"answer": …}}]}',
            "// 항목마다 thought 칸이 먼저 온다",
        ],
        size=11,
    )
    card(
        s,
        X0 + 6.55,
        4.75,
        CW - 6.55,
        1.85,
        [
            {"text": "형식을 틀릴 수 없다", "bold": True, "size": 17, "color": TEAL},
            {
                "text": "디코딩 단계에서 강제하므로 약한 모델도 tool 호출 형식은 항상 맞다. "
                "대가: 응답마다 tool 하나, 생각은 thought 칸에",
                "size": 15,
            },
        ],
        margin=0.25,
        after=6,
    )


def oneshot_slide(d: Deck):
    s = d.content(
        "대화형이 아니라 명령 1회 처리기",
        "상태는 대화가 아니라 작업 폴더의 단계 파일에 있다",
        "대화형으로 만들지 1회용 명령 처리기로 만들지 고민했고, 1회용을 골랐습니다. 이유는 평가 "
        "때문입니다. 같은 명령을 같은 조건으로 다시 돌릴 수 있어야 하고, 실행 하나가 기록 폴더 "
        "하나로 끝나야 비교하기 쉽습니다. 그러면 이어서 일하는 건 어떻게 하느냐. 상태가 대화가 아니라 "
        "작업 폴더의 단계 파일에 남기 때문에, 다음 명령이 stage_status 로 지금 상태를 보고 이어서 "
        "일합니다. 대화형이 필요하면 같은 루프에 입력 반복만 붙이면 됩니다.",
    )
    card(
        s,
        X0,
        1.85,
        3.7,
        2.2,
        [
            {"text": "대화형", "bold": True, "size": 20, "color": GRAY},
            {"text": "상태가 대화 기록에 쌓인다", "bullet": True},
            {"text": "같은 조건으로 다시 돌리기 어렵다", "bullet": True},
        ],
        size=15,
        margin=0.25,
        after=6,
    )
    card(
        s,
        X0,
        4.3,
        3.7,
        2.3,
        [
            {"text": "명령 1회 처리기", "bold": True, "size": 20, "color": TEAL},
            {"text": "명령 하나 → 기록 폴더 하나", "bullet": True},
            {"text": "같은 명령을 다시 돌려 비교", "bullet": True},
            {"text": "상태는 단계 파일에", "bullet": True},
        ],
        fill=TEAL_TINT,
        size=15,
        margin=0.25,
        after=6,
    )
    x = X0 + 4.1
    w = CW - 4.1
    steps = [
        (
            "명령 1",
            '"cone_0019 를 전개도까지 판독해줘"',
            "list_drawings → run_stages(read → develop)",
        ),
        ("결과", "작업 폴더에 stages/01 … 10 이 남는다", "다음 명령이 읽을 상태"),
        (
            "명령 2",
            '"view_labels 부터 다시 돌려줘"',
            "stage_status 로 확인 → run_view_labels → …",
        ),
    ]
    for i, (label, head, sub) in enumerate(steps):
        y = 1.85 + i * 1.62
        circle(s, x, y + 0.25, 0.75, label, fill=ORANGE if i != 1 else TEAL, size=13)
        card(
            s,
            x + 0.95,
            y,
            w - 0.95,
            1.3,
            [
                {"text": head, "bold": True, "size": 17},
                {"text": sub, "size": 14, "color": GRAY, "font": MONO},
            ],
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.25,
            after=4,
        )


def demo_slide(d: Deck):
    s = d.content(
        "시연: cone_0019 판독",
        "로컬 Qwen3.6-27B 가 명령 한 줄을 받아 tool 두 번으로 끝까지 돌렸다 (2026-10-23, mlic)",
        "실제로 돌려 본 결과입니다. 합성 도면 cone_0019 에 '전개도까지 판독하고 결과를 알려줘'라고 "
        "명령했습니다. orchestrator 는 로컬 Qwen3.6-27B, schema 모드였고, 단계 안의 모델은 "
        "DeepSeek-OCR-2 와 같은 Qwen 입니다. 모델은 먼저 list_drawings 로 도면과 작업 폴더를 찾고, "
        "아직 돌린 단계가 없으니 run_stages 로 read 부터 develop 까지 한 번에 돌린 뒤 보고했습니다. "
        "모델 턴은 3번, 전체 133초 중 92초가 단계 실행이었습니다. 보고에 나온 숫자는 전부 tool 결과에 "
        "있던 값이고, 지어낸 숫자는 없었습니다.",
    )
    stats = [
        ("3", "모델 턴"),
        ("133 s", "전체 시간"),
        ("92 s", "그중 단계 실행"),
        ("0", "지어낸 숫자"),
    ]
    w = (CW - 3 * 0.25) / 4
    for i, (value, label) in enumerate(stats):
        x = X0 + i * (w + 0.25)
        card(
            s,
            x,
            1.85,
            w,
            1.75,
            [
                {"text": value, "bold": True, "size": 48, "color": ORANGE},
                {"text": label, "size": 15, "color": GRAY},
            ],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            after=0,
        )
    flow = [
        ("명령", ORANGE_TINT),
        ('list_drawings(name="cone_0019")', LIGHT),
        ("run_stages(read → develop)", LIGHT),
        ("finish: 보고", TEAL_TINT),
    ]
    widths = [1.2, 4.15, 3.6, 1.95]
    x = X0
    for i, ((label, fill), w) in enumerate(zip(flow, widths)):
        chip(s, x, 3.95, w, 0.6, label, fill=fill, size=13, bold=True)
        if i < 3:
            arrow(s, x + w + 0.04, 4.09, 0.25, 0.32)
        x += w + 0.33
    card(
        s,
        X0,
        4.85,
        CW,
        1.0,
        [
            [
                {"text": "턴 2 의 thought  ", "bold": True, "color": TEAL},
                {
                    "text": '"아직 실행된 단계가 없으므로, 요청에 따라 전개도(develop)까지 '
                    '판독하기 위해 read 부터 develop 까지의 단계를 run_stages 로 한 번에 실행합니다."',
                    "italic": True,
                },
            ]
        ],
        size=15,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )
    text(
        s,
        X0,
        6.05,
        CW,
        0.5,
        [
            "orchestrator: Qwen3.6-27B (GPU 1–3, schema 모드) · 단계 모델: DeepSeek-OCR-2 + "
            "Qwen3.6-27B · 마지막 턴 입력 4.3k 토큰"
        ],
        size=13,
        color=GRAY,
    )


def honest_slide(d: Deck):
    s = d.content(
        "된 것과 안 된 것",
        "orchestrator 의 가치는 판단이 필요한 상황에서 나온다",
        "정직하게 평가하겠습니다. 된 것은 자연어 명령을 알맞은 tool 호출로 바꿨고, 보고한 숫자가 모두 "
        "tool 결과에 있었고, 로컬 모델만으로 끝까지 돌았다는 점입니다. 한계도 분명합니다. 처음부터 "
        "끝까지 그대로 돌리는 일에는 이점이 없습니다. 고정 순서와 같은 일을 하고 40초 정도 더 "
        "걸립니다. 판독 정확도도 그대로입니다. 작은 지름이 92 로 나왔는데 정답은 100 이고, 이건 09단계 "
        "판독 오차라 orchestrator 와는 관계없습니다. native 모드, Claude, 여러 도면, 실패 복구는 아직 "
        "확인하지 않았습니다. 그래서 다음 일은 실패나 review 가 생기는 상황에서 고정 순서보다 나은지를 "
        "재는 것입니다.",
    )
    done = [
        "자연어 명령 → 알맞은 tool 호출",
        "보고한 숫자(지름 150 / 92, 높이 150, 반지름 405.65 / 252.87, 68.33°)는 모두 tool 결과에 있다",
        "로컬 모델만으로 끝까지 돌았다",
    ]
    limits = [
        "처음부터 끝까지 돌리기에는 이점 없음 — 고정 순서와 같은 일, 약 +40 s",
        "판독 정확도는 그대로 — 작은 지름 92, 정답 100 (09 unfold_params 오차)",
        "native 모드, Claude, 여러 도면, 실패 복구는 미검증",
    ]
    w = (CW - 0.3) / 2
    for i, (head, items, color, fill) in enumerate(
        [("된 것", done, TEAL, TEAL_TINT), ("한계", limits, ORANGE, ORANGE_TINT)]
    ):
        card(
            s,
            X0 + i * (w + 0.3),
            1.85,
            w,
            3.6,
            [{"text": head, "bold": True, "size": 22, "color": color}]
            + [{"text": item, "bullet": True} for item in items],
            fill=fill,
            size=16,
            margin=0.3,
            after=10,
        )
    card(
        s,
        X0,
        5.7,
        CW,
        0.9,
        [
            [
                {"text": "다음에 잴 것  ", "bold": True, "color": ORANGE},
                {
                    "text": "실패 · review · 부분 재실행이 필요한 도면에서 고정 순서보다 나은가"
                },
            ]
        ],
        size=16,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )


def record_slide(d: Deck):
    s = d.content(
        "모든 판단이 기록으로 남는다",
        "어떤 생각으로 어떤 tool 을 불렀는지가 턴마다 남아, 비교 · 재현 · 평가의 데이터가 된다",
        "하네스에서 기록은 중요한 요소입니다. 실행마다 폴더가 하나 생기고, trajectory.jsonl 에 모델 "
        "턴마다 한 줄, tool 호출마다 한 줄이 남습니다. 모델이 무슨 생각으로 어떤 tool 을 불렀는지, "
        "토큰을 얼마나 썼는지, tool 이 몇 초 걸렸는지가 다 있습니다. 왼쪽이 실제 기록을 줄인 것입니다. "
        "이 기록이 있어야 로컬과 Claude 를 비교하고, 바꾸기 전후를 숫자로 볼 수 있습니다.",
    )
    code(
        s,
        X0,
        1.85,
        7.9,
        4.75,
        [
            "// trajectory.jsonl (cone_0019, 2026-10-23)",
            '{"turn": 1, "seconds": 6.1,',
            ' "text": "… list_drawings를 호출합니다.",',
            ' "calls": [{"name": "list_drawings",',
            '            "arguments": {"name": "cone_0019"}}],',
            ' "usage": {"prompt_tokens": 2318, "completion_tokens": 53}}',
            '{"turn": 1, "tool": "list_drawings", "seconds": 0.0,',
            ' "result": {"count": 1, "drawings": [ … ]}}',
            '{"turn": 2, "calls": [{"name": "run_stages",',
            '   "arguments": {"first": "read", "last": "develop", …}}]}',
            '{"turn": 2, "tool": "run_stages", "seconds": 92.3,',
            ' "result": {"status": "done", "stages": [ … ]}}',
            '{"turn": 3, "text": "cone_0019 도면 판독이 완료되었습니다. …",',
            ' "calls": [], "usage": {"prompt_tokens": 4328}}',
        ],
        size=13,
    )
    files = [
        ("command.json", "명령, 모델, tool 방식, tool 목록 전체"),
        ("trajectory.jsonl", "턴마다 · tool 호출마다 한 줄"),
        ("messages.json", "모델에게 보낸 대화 그대로"),
        ("answer.json", "마지막 보고"),
    ]
    x = X0 + 8.2
    w = CW - 8.2
    for i, (name, what) in enumerate(files):
        card(
            s,
            x,
            1.85 + i * 1.22,
            w,
            1.0,
            [
                {"text": name, "bold": True, "size": 15, "font": MONO, "color": TEAL},
                {"text": what, "size": 13, "color": GRAY},
            ],
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.2,
            after=2,
        )


def next_slide(d: Deck):
    s = d.content(
        "다음 단계",
        "목표는 '로컬 orchestration 이 고정 순서보다 나쁘지 않다'를 숫자로 보이는 것",
        "다음 단계입니다. 첫째, 추론 서버를 tool parser 와 함께 다시 띄워 native 모드를 확인합니다. "
        "둘째, 합성 도면 30장에서 고정 순서와 orchestrator 를 비교합니다. 단계 결과가 같은지, 턴 수, "
        "시간, 토큰을 봅니다. 셋째, 같은 루프로 Claude 와 로컬을 비교합니다. 넷째, 모델 호출이 "
        "실패하거나 review 가 생겼을 때 다시 묻기 같은 보완 장치를 orchestrator 가 판단해서 쓰게 "
        "합니다. 다섯째, v2 의 분할, 네스팅, POR skill 을 tool 로 붙여 판독 뒤까지 잇습니다.",
    )
    steps = [
        ("native 모드 확인", "추론 서버를 tool parser 와 함께 재기동"),
        ("고정 순서 vs orchestrator", "합성 30장: 단계 결과, 턴, 시간, 토큰"),
        ("Claude vs 로컬", "같은 tool · 루프 · 기록으로 비교"),
        ("보완 장치", "모델 오류 · review 가 있으면 다시 묻기를 orchestrator 가 판단"),
        ("뒷단 연결", "v2 skill(분할 · 네스팅 · POR)을 tool 로 추가"),
    ]
    w, gap = 2.3, 0.158
    shape(s, MSO_SHAPE.RECTANGLE, X0 + 0.5, 2.3, CW - 1.0, 0.05, "C9D3D6")
    for i, (head, body) in enumerate(steps):
        x = X0 + i * (w + gap)
        circle(
            s,
            x + w / 2 - 0.4,
            1.95,
            0.8,
            str(i + 1),
            fill=ORANGE if i < 2 else TEAL,
            size=22,
        )
        card(
            s,
            x,
            3.0,
            w,
            2.4,
            [
                {"text": head, "bold": True, "size": 17},
                {"text": body, "size": 15, "color": GRAY},
            ],
            margin=0.22,
            after=8,
        )
    card(
        s,
        X0,
        5.65,
        CW,
        0.9,
        [
            [
                {"text": "기준선  ", "bold": True, "color": ORANGE},
                {
                    "text": "고정 순서 run.py. 바꿀 때마다 합성 30장으로 전후를 비교하고 "
                    "docs/harness_engineering.md 에 적는다"
                },
            ]
        ],
        fill=ORANGE_TINT,
        size=16,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )


def closing_slide(d: Deck):
    s = d.slide(
        "정리하겠습니다. 판독기는 LLM 을 부품으로 쓰는 고정 workflow 였고, 이제 LLM 이 단계를 골라 "
        "부르는 orchestrator 가 생겼습니다. 이를 위해 필요한 것은 MCP 가 아니라 tool calling 이었습니다. "
        "그리고 같은 루프로 로컬 모델과 Claude 를 바꿔 끼울 수 있게 되어, 다음에는 둘을 숫자로 "
        "비교하겠습니다. 질문 받겠습니다.",
        dark=True,
    )
    text(s, X0, 0.8, CW, 0.8, ["정리"], size=36, bold=True, color=WHITE)
    points = [
        "판독기는 고정 workflow 였다 → 이제 LLM 이 단계를 골라 부른다",
        "필요한 것은 MCP 가 아니라 tool calling 이었다",
        "같은 루프로 로컬과 Claude 를 비교할 수 있다 → 다음은 숫자로 보이기",
    ]
    for i, point in enumerate(points):
        y = 2.0 + i * 1.25
        circle(s, X0, y, 0.7, str(i + 1), size=20)
        text(
            s,
            X0 + 1.0,
            y,
            CW - 1.0,
            0.7,
            [point],
            size=22,
            color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    text(s, X0, 6.1, CW, 0.6, ["질문"], size=26, bold=True, color=ORANGE)


def main() -> None:
    d = Deck()
    for build in (
        title_slide,
        goal_slide,
        reader_slide,
        toolcall_slide,
        mcp_slide,
        architecture_slide,
        tools_slide,
        modes_slide,
        oneshot_slide,
        demo_slide,
        honest_slide,
        record_slide,
        next_slide,
        closing_slide,
    ):
        build(d)
    print(f"{d.count} slides -> {d.save()}")


if __name__ == "__main__":
    main()
