"""Build presentation/dwg_dxf.pptx: DWG and DXF, why the project converts, and how.

PYTHONPATH=. uv run --with python-pptx python presentation/build_dwg_dxf.py

Facts checked on mlic on 2026-10-23 (docs/presentation.md 4절): the synthetic cone_0019.dxf;
the 536 input DWGs under Data/dataset/KDM (versions and sizes only; customer drawings, so
no content appears on the slides); 11 of them converted with LibreDWG 0.14 dwg2dxf into the
session scratchpad, three measured with the Qwen3.6-27B tokenizer and ezdxf. Conversion
failures over all 536 come from AGENTS.md 6절 (2026-10-22).
"""

from pathlib import Path

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from presentation.slides import (
    CW,
    GRAY,
    INK,
    LIGHT,
    ORANGE,
    ORANGE_TINT,
    TEAL,
    TEAL_TINT,
    WHITE,
    X0,
    Deck,
    arrow,
    card,
    chip,
    circle,
    code,
    rgb,
    text,
    write,
)

OUT = Path(__file__).with_name("dwg_dxf.pptx")
TITLE = "DWG 와 DXF: 도면 파일을 LLM agent 가 다루려면"
FOOTER = "PBL cutplan_agent · DWG 와 DXF · 2026-10-24"


def stat_row(s, y, stats, h=1.6, value_size=40):
    """Big numbers with a label under each, side by side."""
    w = (CW - (len(stats) - 1) * 0.25) / len(stats)
    for i, (value, label) in enumerate(stats):
        card(
            s,
            X0 + i * (w + 0.25),
            y,
            w,
            h,
            [
                {"text": value, "bold": True, "size": value_size, "color": ORANGE},
                {"text": label, "size": 14, "color": GRAY},
            ],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            after=2,
            margin=0.12,
        )


def callout(s, y, label, body, h=0.95, fill=ORANGE_TINT, color=ORANGE):
    card(
        s,
        X0,
        y,
        CW,
        h,
        [[{"text": label + "  ", "bold": True, "color": color}, {"text": body}]],
        fill=fill,
        size=16,
        anchor=MSO_ANCHOR.MIDDLE,
        margin=0.3,
    )


def title_slide(d: Deck):
    s = d.slide(
        "이번 파트는 도면 파일 형식 이야기입니다. 고객 도면은 대부분 DWG 로 오는데, 우리 파이프라인은 "
        "전부 DXF 를 읽습니다. 두 형식이 무엇이 다른지, 왜 DXF 로 바꾸는지, 그리고 'DXF 는 텍스트라서 "
        "LLM 이 읽을 수 있다'는 말이 사실인지 실제 파일로 확인한 결과를 보여 드리고, 마지막으로 "
        "변환하는 방법과 그 한계를 정리하겠습니다.",
        dark=True,
    )
    text(
        s,
        X0,
        2.2,
        CW,
        1.6,
        ["DWG 와 DXF", "도면 파일을 LLM agent 가 다루려면"],
        size=40,
        bold=True,
        color=WHITE,
        after=4,
    )
    text(
        s,
        X0,
        4.0,
        CW,
        0.5,
        ["두 형식의 차이, DXF 로 바꾸는 이유, 변환 방법과 한계"],
        size=20,
        color=ORANGE,
    )
    text(
        s,
        X0,
        6.4,
        CW,
        0.4,
        ["PBL cutplan_agent · 2026-10-24"],
        size=14,
        color="B8C4C9",
    )


def compare_slide(d: Deck):
    s = d.content(
        "DWG 와 DXF 한눈에",
        "같은 도면을 담지만, 하나는 AutoCAD 의 저장 형식이고 하나는 교환용 공개 형식이다",
        "두 형식을 나란히 놓았습니다. DWG 는 AutoCAD 가 기본으로 저장하는 형식이고 바이너리입니다. "
        "Autodesk 는 DWG 의 공식 사양을 공개하지 않았고, Open Design Alliance 가 역공학으로 만든 "
        "사양과 SDK, 그리고 GNU LibreDWG 같은 오픈소스가 읽습니다. DXF 는 Autodesk 가 다른 "
        "프로그램과 도면을 주고받으라고 사양을 공개한 형식이고, 보통 텍스트로 저장됩니다. 바이너리 "
        "DXF 도 있지만 드뭅니다. 우리 데이터로 재 보니 같은 도면이 DXF 로 바꾸면 5배에서 8배 "
        "커졌습니다.",
    )
    rows = [
        ("누가 만들었나", "AutoCAD 의 기본 저장 형식", "Autodesk 가 공개한 교환 형식"),
        ("저장 방식", "바이너리 (압축)", "텍스트 (바이너리 DXF 도 있지만 드묾)"),
        (
            "사양",
            "Autodesk 공식 사양 없음. ODA 가 역공학",
            "Autodesk DXF Reference 로 공개",
        ),
        (
            "읽는 도구",
            "AutoCAD, ODA SDK, LibreDWG",
            "ezdxf 등 대부분의 CAD 라이브러리",
        ),
        ("크기 (우리 도면 9장)", "0.2–0.5 MB", "1.1–2.8 MB (DWG 의 5.0–7.7배)"),
    ]
    table = s.shapes.add_table(
        len(rows) + 1, 3, Inches(X0), Inches(1.85), Inches(CW), Inches(4.0)
    ).table
    for c, w in enumerate([2.9, 4.5, 4.733]):
        table.columns[c].width = Inches(w)
    header = ["", "DWG", "DXF"]
    for r, row in enumerate([header, *rows]):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(INK)
            elif c == 2:
                cell.fill.fore_color.rgb = rgb(TEAL_TINT)
            else:
                cell.fill.fore_color.rgb = rgb(LIGHT)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            write(
                cell.text_frame,
                [value],
                size=18 if r == 0 else 15,
                bold=r == 0 or c == 0,
                color=WHITE if r == 0 else (GRAY if c == 0 else INK),
                after=0,
                margin=0.12,
            )
    text(
        s,
        X0,
        6.05,
        CW,
        0.5,
        [
            "크기: Data/dataset/KDM 입력 도면 DWG 11장을 LibreDWG 0.14 로 변환, 성공한 9장 "
            "(2026-10-23, mlic)"
        ],
        size=12,
        color=GRAY,
    )


def inside_slide(d: Deck):
    s = d.content(
        "파일을 열어 보면",
        "DWG 는 버전 문자열 뒤가 바이너리이고, DXF 는 '그룹 코드 / 값' 두 줄이 반복되는 텍스트다",
        "실제로 열어 보면 차이가 분명합니다. DWG 는 처음 6바이트가 AC1018 같은 버전 문자열이고, "
        "그 뒤는 압축된 바이너리라 사람이 읽을 수 없습니다. 오른쪽은 합성 도면 cone_0019.dxf 에서 선 "
        "하나를 그대로 옮긴 것입니다. 한 줄은 그룹 코드, 다음 줄은 값입니다. 0 은 엔티티 시작, 8 은 "
        "레이어, 10/20/30 은 시작점, 11/21/31 은 끝점입니다. 텍스트라서 메모장으로도 열리고 grep 도 "
        "됩니다. 다만 '이 선이 정면도의 윗변'이라는 뜻은 어디에도 없고 좌표만 있습니다.",
    )
    card(
        s,
        X0,
        1.85,
        5.4,
        4.75,
        [
            {"text": "DWG", "bold": True, "size": 24, "color": GRAY},
            {"text": "처음 6바이트 = 버전 문자열", "bold": True, "size": 16},
            {
                "text": "그 뒤는 압축된 바이너리. 텍스트 편집기로 열면 깨진 글자만 보인다",
                "size": 15,
            },
            {"text": "우리 입력 도면 536장의 버전", "bold": True, "size": 16},
        ],
        margin=0.3,
        after=10,
    )
    versions = [
        ("AC1021", "R2007 형식", "301장"),
        ("AC1018", "R2004 형식", "231장"),
        ("AC1024", "R2010 형식", "4장"),
    ]
    for i, (magic, name, count) in enumerate(versions):
        y = 4.55 + i * 0.62
        chip(
            s, X0 + 0.3, y, 1.5, 0.48, magic, fill=INK, color=WHITE, size=14, bold=True
        )
        text(
            s,
            X0 + 2.0,
            y + 0.06,
            3.2,
            0.4,
            [[{"text": name + "  "}, {"text": count, "bold": True, "color": ORANGE}]],
            size=15,
        )
    code(
        s,
        X0 + 5.7,
        1.85,
        CW - 5.7,
        4.75,
        [
            "// cone_0019.dxf 의 LINE 하나 (그대로)",
            "0          ← 0: 엔티티 시작",
            "LINE",
            "8          ← 8: 레이어",
            "DIM",
            "10         ← 10 / 20 / 30: 시작점 x, y, z",
            "100.0",
            "20",
            "86.0",
            "30",
            "0",
            "11         ← 11 / 21 / 31: 끝점",
            "250.0",
            "21",
            "86.0",
            "31",
            "0",
        ],
        size=12,
    )


def why_slide(d: Deck):
    s = d.content(
        "왜 DXF 로 만드나",
        "고객 도면은 DWG 로 오지만, 우리 파이프라인은 전부 DXF 를 읽는다",
        "그래서 DXF 로 바꿉니다. 이유는 네 가지입니다. 첫째, 우리가 쓰는 파이썬 라이브러리 ezdxf 가 "
        "DWG 를 읽지 못합니다. 판독기, 도면 렌더링 스크립트, v2 에이전트가 모두 DXF 를 읽습니다. "
        "둘째, DXF 는 사양이 공개돼 있어서 오픈소스로 끝까지 다룰 수 있습니다. 셋째, 텍스트라서 사람이 "
        "열어 보고 grep 이나 diff 로 비교할 수 있어 디버깅이 쉽습니다. 넷째, 동료의 원래 파이프라인도 "
        "툴킷의 dwg2dxf.exe 로 DXF 를 만들어 썼습니다. 우리는 같은 LibreDWG 0.14 를 리눅스에 설치해 "
        "씁니다.",
    )
    flow = [
        ("DWG", "고객 도면", LIGHT),
        ("dwg2dxf", "LibreDWG 0.14", ORANGE_TINT),
        ("DXF", "텍스트", TEAL_TINT),
        ("ezdxf", "파이썬으로 파싱", LIGHT),
        ("판독 · 렌더링", "reader, render_drawings, v2", LIGHT),
    ]
    w, gap = 2.15, (CW - 5 * 2.15) / 4
    for i, (head, sub, fill) in enumerate(flow):
        x = X0 + i * (w + gap)
        card(
            s,
            x,
            1.85,
            w,
            1.05,
            [
                {"text": head, "bold": True, "size": 17},
                {"text": sub, "size": 12, "color": GRAY},
            ],
            fill=fill,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.08,
            after=2,
        )
        if i < 4:
            arrow(s, x + w + (gap - 0.3) / 2, 2.22, 0.3, 0.32)
    reasons = [
        (
            "ezdxf 는 DWG 를 읽지 못한다",
            "판독기, render_drawings.py, v2 ingest_drawing 이 모두 DXF 를 읽는다",
        ),
        ("사양이 공개돼 있다", "DXF Reference 가 있어 오픈소스로 끝까지 다룰 수 있다"),
        (
            "사람이 열어 볼 수 있다",
            "텍스트라 grep · diff 로 찾고 비교한다. 디버깅이 쉽다",
        ),
        (
            "동료 파이프라인과 같다",
            "툴킷의 dwg2dxf.exe 와 같은 LibreDWG 0.14 를 리눅스에 설치해 쓴다",
        ),
    ]
    w2 = (CW - 0.25) / 2
    for i, (head, body) in enumerate(reasons):
        row, col = divmod(i, 2)
        x = X0 + col * (w2 + 0.25)
        y = 3.2 + row * 1.7
        circle(s, x + 0.2, y + 0.3, 0.6, str(i + 1), size=18)
        card(
            s,
            x + 1.0,
            y,
            w2 - 1.0,
            1.45,
            [
                {"text": head, "bold": True, "size": 17, "color": TEAL},
                {"text": body, "size": 14},
            ],
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.22,
            after=4,
        )


def factcheck_slide(d: Deck):
    s = d.content(
        "사실 확인: DXF 는 LLM 이 읽을 수 있는 텍스트?",
        "텍스트인 것은 맞다. 하지만 실제 도면은 그대로 넣기에는 너무 크다",
        "'DXF 는 텍스트라서 LLM agent 가 읽을 수 있다'는 말을 실제 파일로 확인했습니다. 텍스트인 것은 "
        "맞습니다. 그런데 토큰 수를 Qwen3.6-27B 토크나이저로 세어 보니, 아주 단순한 합성 도면도 "
        "4,700 토큰이고, 실제 입력 도면을 DXF 로 바꾼 것은 82만에서 218만 토큰이었습니다. 지금 띄워 둔 "
        "서버의 컨텍스트는 16,384 토큰이고, 모델 자체의 최대치도 262,144 토큰입니다. 실제 도면은 서버 "
        "한도의 50배 이상, 모델 최대치의 3배 이상이라 통째로 넣을 수 없습니다. 그래서 결론은 '반만 "
        "맞다'입니다.",
    )
    stat_row(
        s,
        1.85,
        [
            ("4,747", "합성 cone_0019.dxf\n5.5 KB, 엔티티 66개"),
            ("82–218만", "실제 입력 도면 DXF 3장\n1.1–2.8 MB"),
            ("16,384", "지금 vLLM 서버의\nmax-model-len"),
            ("262,144", "Qwen3.6 토크나이저의\n최대 길이"),
        ],
        h=2.0,
        value_size=32,
    )
    text(
        s,
        X0,
        3.95,
        CW,
        0.4,
        ["토큰 수: Qwen3.6-27B 토크나이저로 DXF 파일 전체를 셈 (2026-10-23)"],
        size=12,
        color=GRAY,
    )
    w = (CW - 0.3) / 2
    card(
        s,
        X0,
        4.5,
        w,
        1.4,
        [
            {"text": "맞는 부분", "bold": True, "size": 18, "color": TEAL},
            {
                "text": "ASCII 텍스트다. 메모장 · grep · diff 로 열리고, 작은 도면은 통째로 넣을 수 있다",
                "size": 15,
            },
        ],
        fill=TEAL_TINT,
        margin=0.25,
        after=4,
    )
    card(
        s,
        X0 + w + 0.3,
        4.5,
        w,
        1.4,
        [
            {"text": "틀린 부분", "bold": True, "size": 18, "color": ORANGE},
            {
                "text": "실제 도면은 서버 한도의 50배 이상, 모델 최대치의 3배 이상. 통째로는 못 넣는다",
                "size": 15,
            },
        ],
        fill=ORANGE_TINT,
        margin=0.25,
        after=4,
    )
    text(
        s,
        X0,
        6.1,
        CW,
        0.5,
        [
            [
                {"text": "판정  ", "bold": True, "color": ORANGE},
                {
                    "text": "반만 맞다. 크기 말고도 걸리는 것이 두 가지 더 있다 (다음 장)"
                },
            ]
        ],
        size=16,
    )


def barriers_slide(d: Deck):
    s = d.content(
        "그대로 넣으면 안 되는 세 가지 이유",
        "크기, 의미, 인코딩",
        "크기 말고도 두 가지가 더 있습니다. 둘째는 의미입니다. DXF 에는 숫자 그룹 코드와 좌표만 "
        "있습니다. 이 선이 정면도에 속한다거나, 이 치수가 지름 150 이라는 뜻은 엔티티 사이의 위치 "
        "관계에서 나옵니다. 실제 도면 하나에 엔티티가 2천에서 7천 개이고 그중 60~70% 가 선입니다. "
        "LLM 에게 수천 개 좌표 사이의 관계를 계산시키는 것은 '계산은 코드가 한다'는 우리 원칙에도 "
        "어긋납니다. 셋째는 인코딩입니다. 한글이 CP949 로 저장된 파일이 있어서 UTF-8 로 읽으면 "
        "깨집니다. 샘플에서 62자가 깨졌고, CP949 로 읽으면 한글 39자가 제대로 나왔습니다.",
    )
    items = [
        (
            "1",
            "크기",
            "실제 도면 DXF 는 82만–218만 토큰. 서버 한도(16,384)의 50배 이상",
            "1.1–2.8 MB / 12만–34만 줄",
        ),
        (
            "2",
            "의미",
            "그룹 코드와 좌표뿐이다. '정면도', '지름 150' 같은 뜻은 엔티티 사이의 위치 관계에서 나온다",
            "엔티티 2천–7천 개, LINE 60–69%",
        ),
        (
            "3",
            "인코딩",
            "한글이 CP949(ANSI_949)로 저장된 파일이 있다. UTF-8 로 읽으면 깨진다",
            "UTF-8 62자 깨짐 / CP949 한글 39자",
        ),
    ]
    w = (CW - 2 * 0.25) / 3
    for i, (num, head, body, fact) in enumerate(items):
        x = X0 + i * (w + 0.25)
        card(
            s,
            x,
            1.85,
            w,
            4.0,
            [
                {"text": head, "bold": True, "size": 22, "color": TEAL},
                {"text": body, "size": 16},
            ],
            margin=0.28,
            after=10,
        ).text_frame.margin_top = Inches(1.05)
        circle(s, x + 0.28, 2.05, 0.75, num, size=22)
        chip(s, x + 0.2, 5.05, w - 0.4, 0.55, fact, fill=WHITE, color=GRAY, size=12)
    text(
        s,
        X0,
        6.05,
        CW,
        0.5,
        [
            "엔티티 수 · 인코딩: 변환한 입력 도면 DXF 를 ezdxf 로 읽어 셈. 도면 내용은 싣지 않는다 (고객 도면)"
        ],
        size=12,
        color=GRAY,
    )


def pipeline_slide(d: Deck):
    s = d.content(
        "그래서 우리는: 코드가 읽고, LLM 은 좁게 본다",
        "DXF 가 텍스트라서 생기는 이점은 LLM 보다 코드와 사람 쪽에 있다",
        "그래서 판독기는 DXF 를 LLM 에 통째로 주지 않습니다. ezdxf 로 코드가 파싱하고, 코드가 뷰를 "
        "나누고 치수와 표제란을 찾습니다. LLM 에게는 그린 그림, 표제란을 OCR 한 글자, 코드가 정리한 "
        "요약만 주고, 뷰 종류나 형상처럼 좁은 판단만 맡깁니다. agent 로 확장할 때도 같습니다. DXF "
        "전체를 넣는 대신, 표제란 텍스트나 특정 레이어처럼 필요한 부분만 꺼내 주는 tool 을 만드는 "
        "것이 맞습니다. DXF 가 텍스트라서 좋은 점은 오히려 코드와 사람이 쉽게 읽는다는 데 있습니다.",
    )
    steps = [
        ("DXF", "텍스트 파일", LIGHT, INK),
        ("코드: ezdxf", "엔티티 66 – 7,245개 파싱", TEAL, WHITE),
        ("코드: 판독 단계", "뷰 군집 · 치수 · 표제란 위치", TEAL, WHITE),
        ("LLM 입력", "그림 · 표제란 OCR 글자 · 요약", ORANGE, WHITE),
        ("LLM 판단", "뷰 종류 · 형상 · 전개 파라미터", ORANGE_TINT, INK),
    ]
    w, gap = 2.15, (CW - 5 * 2.15) / 4
    for i, (head, sub, fill, color) in enumerate(steps):
        x = X0 + i * (w + gap)
        s_ = card(
            s,
            x,
            1.95,
            w,
            1.5,
            [
                {"text": head, "bold": True, "size": 17, "color": color},
                {"text": sub, "size": 13, "color": color},
            ],
            fill=fill,
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.1,
            after=4,
        )
        s_.name = f"step {i + 1}"
        if i < 4:
            arrow(s, x + w + (gap - 0.3) / 2, 2.54, 0.3, 0.32)
    w2 = (CW - 0.3) / 2
    card(
        s,
        X0,
        3.85,
        w2,
        2.0,
        [
            {
                "text": "텍스트의 이점은 코드와 사람에게",
                "bold": True,
                "size": 18,
                "color": TEAL,
            },
            {"text": "오픈소스로 파싱", "bullet": True},
            {"text": "grep · diff 로 찾고 비교", "bullet": True},
            {"text": "변환이 잘못되면 열어서 확인", "bullet": True},
        ],
        fill=TEAL_TINT,
        size=15,
        margin=0.25,
        after=4,
    )
    card(
        s,
        X0 + w2 + 0.3,
        3.85,
        w2,
        2.0,
        [
            {"text": "agent 로 확장한다면", "bold": True, "size": 18, "color": ORANGE},
            {
                "text": "DXF 전체 대신 필요한 부분만 꺼내는 tool 을 준다 "
                "(예: 표제란 텍스트, 특정 레이어의 TEXT, 한 뷰의 치수)",
                "size": 15,
            },
        ],
        fill=ORANGE_TINT,
        margin=0.25,
        after=6,
    )
    text(
        s,
        X0,
        6.05,
        CW,
        0.5,
        ["판독 단계의 자세한 구성은 orchestrator 파트와 docs/reader/stages.md"],
        size=13,
        color=GRAY,
    )


def tools_slide(d: Deck):
    s = d.content(
        "DWG 를 DXF 로 바꾸는 도구",
        "이 레포는 오픈소스 LibreDWG 를 쓴다. 실패하는 도면은 ODA File Converter 가 다음 후보",
        "DWG 를 DXF 로 바꾸는 방법은 크게 세 가지입니다. 첫째는 GNU LibreDWG 의 dwg2dxf 로, 오픈소스 "
        "명령행 도구이고 이 레포가 쓰는 방법입니다. 둘째는 Open Design Alliance 의 ODA File Converter "
        "입니다. 무료이고 윈도우, 리눅스, 맥에서 GUI 와 명령행 모두 되며, ezdxf 에도 이 도구를 부르는 "
        "odafc 애드온이 있습니다. 이 서버에는 아직 설치하지 않았습니다. 셋째는 AutoCAD 에서 다른 "
        "이름으로 저장할 때 DXF 를 고르거나 DXFOUT 명령을 쓰는 방법인데, 상용 라이선스가 필요합니다.",
    )
    rows = [
        ["도구", "종류", "플랫폼", "이 레포에서"],
        [
            "LibreDWG dwg2dxf",
            "오픈소스 (GNU, GPL)\n명령행",
            "Linux, Windows",
            "사용 중 (0.14)\nscripts/install_libredwg.sh",
        ],
        [
            "ODA File Converter",
            "무료 (Open Design Alliance)\nGUI + 명령행",
            "Windows, Linux, macOS",
            "미설치\nezdxf.addons.odafc 로 부를 수 있다",
        ],
        [
            "AutoCAD",
            "상용\n다른 이름으로 저장 → DXF, DXFOUT",
            "Windows, macOS",
            "쓰지 않음",
        ],
    ]
    table = s.shapes.add_table(
        len(rows), 4, Inches(X0), Inches(1.85), Inches(CW), Inches(3.6)
    ).table
    for c, w in enumerate([2.7, 3.4, 2.4, 3.633]):
        table.columns[c].width = Inches(w)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(INK)
            elif r == 1:
                cell.fill.fore_color.rgb = rgb(TEAL_TINT)
            else:
                cell.fill.fore_color.rgb = rgb(LIGHT if r % 2 == 0 else "F7F9FA")
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            write(
                cell.text_frame,
                [{"text": line} for line in value.split("\n")],
                size=15,
                bold=r == 0 or c == 0,
                color=WHITE if r == 0 else INK,
                after=0,
                margin=0.12,
            )
    callout(
        s,
        5.7,
        "주의",
        "ezdxf 는 DWG 를 직접 읽지 못한다. DWG 는 먼저 변환하거나 odafc 로 ODA File Converter 를 거친다",
        h=0.9,
    )


def howto_slide(d: Deck):
    s = d.content(
        "이 레포에서 변환하기",
        "리눅스는 LibreDWG 를 소스에서 빌드해 ~/.local 에 설치한다 (Ubuntu apt 에는 없다)",
        "이 레포에서 실제로 쓰는 명령입니다. 우분투 apt 에는 LibreDWG 가 없어서, install_libredwg.sh "
        "가 GNU 에서 0.14 소스를 받아 sha256 을 확인하고 sudo 없이 홈 폴더에 빌드합니다. 그다음 "
        "dwg2dxf 에 출력 파일을 주면 됩니다. 사람이 볼 PDF 나 PNG 가 필요하면 render_drawings.py 가 "
        "변환부터 그리기까지 한 번에 합니다. 윈도우는 툴킷에 있는 dwg2dxf.exe 를 씁니다. ODA File "
        "Converter 를 설치하면 아래 명령이나 ezdxf 의 odafc 로도 바꿀 수 있습니다.",
    )
    code(
        s,
        X0,
        1.85,
        7.3,
        4.75,
        [
            "# 1. LibreDWG 0.14 설치 (한 번, sudo 불필요)",
            "scripts/install_libredwg.sh       # → ~/.local/bin/dwg2dxf",
            "",
            "# 2. DWG → DXF",
            "dwg2dxf -y -o out/drawing.dxf in/drawing.dwg",
            "",
            "# 3. 사람이 볼 PDF · PNG (변환 + 그리기)",
            "uv run python scripts/render_drawings.py <DWG 폴더>",
            "",
            "# ODA File Converter 를 설치했다면",
            'ODAFileConverter in/ out/ ACAD2018 DXF 0 1 "*.DWG"',
            "",
            "# 파이썬에서 (ODA File Converter 필요)",
            "from ezdxf.addons import odafc",
            'doc = odafc.readfile("in/drawing.dwg")',
        ],
        size=13,
    )
    notes = [
        ("-y", "있으면 덮어쓰기"),
        ("Windows", "툴킷 tools/dwg2png/libredwg/dwg2dxf.exe (같은 0.14)"),
        (
            "ODA 인자",
            "입력 폴더, 출력 폴더, 출력 버전, 형식, 하위 폴더(0/1), 검사(0/1), 파일 필터",
        ),
        ("v2 에이전트", "ingest_drawing 이 같은 dwg2dxf 를 부른다"),
    ]
    x = X0 + 7.6
    w = CW - 7.6
    for i, (head, body) in enumerate(notes):
        card(
            s,
            x,
            1.85 + i * 1.22,
            w,
            1.05,
            [
                {"text": head, "bold": True, "size": 15, "color": TEAL},
                {"text": body, "size": 13, "color": GRAY},
            ],
            anchor=MSO_ANCHOR.MIDDLE,
            margin=0.2,
            after=2,
        )


def limits_slide(d: Deck):
    s = d.content(
        "변환의 한계: 성공은 결과로 확인한다",
        "LibreDWG 0.14 는 입력 도면 536장 중 87장을 제대로 바꾸지 못했다 (2026-10-22)",
        "변환에는 한계가 있습니다. 입력 도면 DWG 536장을 LibreDWG 0.14 로 바꿔 보니 87장, 약 16% 가 "
        "실패했습니다. 그중 77장은 R2007 형식 DWG 에서 오류를 내고도 종료 코드 0 으로 빈 DXF 를 "
        "썼습니다. 그러니까 명령이 성공했다고 믿으면 안 됩니다. 나머지 10장은 ezdxf 가 열지 못하는 "
        "DXF 를 만들었습니다. 오늘 샘플 11장으로 다시 해 봐도 R2007 2장이 7,463 바이트짜리 빈 DXF 로 "
        "나왔습니다. 교훈은 변환이 됐는지를 종료 코드가 아니라 결과, 예를 들어 엔티티 수로 확인해야 "
        "한다는 것입니다. 실패한 87장 중 47장은 같은 이름의 PDF 가 있고, 나머지는 ODA File Converter "
        "로 다시 시도해 볼 후보입니다.",
    )
    stat_row(
        s,
        1.85,
        [
            ("536", "입력 도면 DWG\n(Data/dataset/KDM)"),
            ("87", "제대로 변환 안 됨\n(약 16%)"),
            ("77", "R2007 에서 빈 DXF\n그런데 종료 코드 0"),
            ("10", "ezdxf 가 열지 못하는\nDXF"),
        ],
        h=1.9,
        value_size=44,
    )
    w = (CW - 0.3) / 2
    card(
        s,
        X0,
        4.0,
        w,
        1.75,
        [
            {
                "text": "오늘 다시 확인 (샘플 11장)",
                "bold": True,
                "size": 17,
                "color": TEAL,
            },
            {
                "text": "R2007 2장이 7,463 바이트짜리 빈 DXF. 나머지 9장은 1.1–2.8 MB",
                "size": 15,
            },
        ],
        margin=0.25,
        after=6,
    )
    card(
        s,
        X0 + w + 0.3,
        4.0,
        w,
        1.75,
        [
            {"text": "실패한 87장은", "bold": True, "size": 17, "color": TEAL},
            {
                "text": "47장은 같은 이름의 PDF 가 옆에 있다. 나머지는 ODA File Converter 로 다시 시도할 후보",
                "size": 15,
            },
        ],
        margin=0.25,
        after=6,
    )
    callout(
        s,
        5.95,
        "교훈",
        "변환 성공은 종료 코드가 아니라 결과(파일 크기, 엔티티 수)로 확인한다",
        h=0.7,
    )


def closing_slide(d: Deck):
    s = d.slide(
        "정리하겠습니다. DWG 는 AutoCAD 의 바이너리 저장 형식이고, DXF 는 공개된 텍스트 교환 형식입니다. "
        "우리 도구가 DXF 를 읽기 때문에 변환합니다. DXF 가 텍스트인 것은 맞지만 실제 도면은 LLM 에 "
        "통째로 넣기엔 너무 크고, 의미와 인코딩 문제도 있어서, 코드가 읽고 LLM 에는 필요한 부분만 "
        "줍니다. 변환은 LibreDWG 로 하고, 성공 여부는 결과로 확인합니다. 질문 받겠습니다.",
        dark=True,
    )
    text(s, X0, 0.8, CW, 0.8, ["정리"], size=36, bold=True, color=WHITE)
    points = [
        "DWG 는 AutoCAD 의 바이너리 저장 형식, DXF 는 공개된 텍스트 교환 형식이다",
        "우리 도구(ezdxf)가 DXF 를 읽기 때문에 DWG 를 DXF 로 바꾼다",
        "DXF 는 텍스트지만 실제 도면은 통째로 넣기엔 너무 크다 → 코드가 읽는다",
        "변환은 LibreDWG dwg2dxf. 성공은 종료 코드가 아니라 결과로 확인한다",
    ]
    for i, point in enumerate(points):
        y = 1.9 + i * 1.05
        circle(s, X0, y, 0.65, str(i + 1), size=18)
        text(
            s,
            X0 + 0.95,
            y,
            CW - 0.95,
            0.65,
            [point],
            size=20,
            color=WHITE,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    text(s, X0, 6.3, CW, 0.6, ["질문"], size=26, bold=True, color=ORANGE)


def main() -> None:
    d = Deck(OUT, TITLE, FOOTER)
    for build in (
        title_slide,
        compare_slide,
        inside_slide,
        why_slide,
        factcheck_slide,
        barriers_slide,
        pipeline_slide,
        tools_slide,
        howto_slide,
        limits_slide,
        closing_slide,
    ):
        build(d)
    print(f"{d.count} slides -> {d.save()}")


if __name__ == "__main__":
    main()
